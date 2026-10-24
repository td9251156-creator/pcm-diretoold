"""Verificações dos arquivos GERADOS (não do código-fonte). Uso, na raiz do repositório:

    python3 entrega-templates/testes/verificar_entrega.py

Requer LibreOffice (soffice) com Python UNO, pdftotext (poppler), openpyxl e python-docx.
Grava testes/resultado-testes.json e as prévias PDF em previas/.
"""
import copy
import datetime as dt
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
sys.path.insert(0, str(AQUI))
sys.path.insert(0, str(RAIZ / "gerador"))

from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import especificacao as E  # noqa: E402
import lo_util as LO  # noqa: E402
from comum import AVISO  # noqa: E402

BUILD = RAIZ / "build" / "testes"
BUILD.mkdir(parents=True, exist_ok=True)
PREVIAS = RAIZ / "previas"
PREVIAS.mkdir(exist_ok=True)
PLAN = RAIZ / "planilhas"
ARQ = {a.chave: a for a in E.ARQUIVOS}
MODELO = {k: PLAN / f"{a.nome}.xlsx" for k, a in ARQ.items()}
EXEMPLO = {k: PLAN / f"{a.nome}-Exemplo.xlsx" for k, a in ARQ.items()}
MAN = RAIZ / "manuais"
DOCX = {
    "W1": MAN / "PCMD-W1-Manual-Gestao-Ativos.docx", "W2": MAN / "PCMD-W2-Manual-Estrategia-Manutencao.docx",
    "W1-Exemplo": MAN / "PCMD-W1-Manual-Gestao-Ativos-Exemplo.docx", "W2-Exemplo": MAN / "PCMD-W2-Manual-Estrategia-Manutencao-Exemplo.docx",
}
RESULTADOS = []


def check(grupo, nome, ok, detalhe=""):
    RESULTADOS.append({"grupo": grupo, "verificacao": nome, "resultado": "PASSOU" if ok else "FALHOU", "detalhe": str(detalhe)[:600]})
    print(f"[{'OK ' if ok else 'FALHA'}] {grupo} · {nome}" + (f" — {detalhe}" if detalhe and not ok else ""))
    return ok


def norm(t):
    t = t.replace("­", "")
    t = re.sub(r"-\s*\n\s*", "-", t)
    return re.sub(r"\s+", " ", t)


# ------------------------------------------------------------------ 1 e 2: abertura e estrutura
def estrutura_zip(caminho, grupo):
    with zipfile.ZipFile(caminho) as z:
        nomes = z.namelist()
        check(grupo, f"{caminho.name}: ZIP íntegro (testzip)", z.testzip() is None)
        check(grupo, f"{caminho.name}: sem macro (vbaProject.bin)", not any("vbaProject" in n for n in nomes))
        check(grupo, f"{caminho.name}: sem vínculos externos (externalLinks)", not any("externalLink" in n for n in nomes))
        check(grupo, f"{caminho.name}: sem conexões de dados", not any("connections.xml" in n or "queryTable" in n for n in nomes))
        rels = "".join(z.read(n).decode("utf-8", "ignore") for n in nomes if n.endswith(".rels"))
        check(grupo, f"{caminho.name}: nenhuma relação TargetMode=External", 'TargetMode="External"' not in rels)
        ct = z.read("[Content_Types].xml").decode()
        check(grupo, f"{caminho.name}: tipo de conteúdo sem macroEnabled", "macroEnabled" not in ct)
        xml_total = "".join(z.read(n).decode("utf-8", "ignore") for n in nomes if n.endswith(".xml"))
        return nomes, xml_total


PROIBIDAS = re.compile(r"\b(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SORTBY|SEQUENCE|LET|LAMBDA|RANDARRAY|TEXTJOIN|IFS|SWITCH|_xlfn|_xlws)\b", re.I)


def verificar_xlsx_estrutura(caminho):
    g = "2 · Estrutura XLSX"
    nomes, xml = estrutura_zip(caminho, g)
    formulas = re.findall(r"<f>(.*?)</f>", xml)
    proib = [f for f in formulas if PROIBIDAS.search(f)]
    check(g, f"{caminho.name}: nenhuma fórmula proibida ({len(formulas)} fórmulas)", not proib, proib[:3])
    check(g, f"{caminho.name}: sem fórmula de matriz dinâmica/CSE", 't="array"' not in xml and "cm=" not in xml)
    wb = load_workbook(caminho)
    ruins = [n for n in wb.sheetnames if len(n) > 31 or re.search(r"[\[\]:*?/\\]", n)]
    check(g, f"{caminho.name}: nomes de abas válidos (≤31, sem []:*?/\\)", not ruins, ruins)
    ruins = [n for n in wb.defined_names if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", n) or re.fullmatch(r"[A-Za-z]{1,3}\d+", n)]
    check(g, f"{caminho.name}: intervalos nomeados válidos ({len(wb.defined_names)})", not ruins, ruins)
    dv_ruins = []
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            for f in (dv.formula1, dv.formula2):
                if f and "!" in f:
                    dv_ruins.append((ws.title, f))
    check(g, f"{caminho.name}: validações sem referência direta a outra aba", not dv_ruins, dv_ruins[:3])
    rod_ruins, aviso_ruins = [], []
    for ws in wb.worksheets:
        bruto = z_rodape(caminho, ws.title)
        if len(bruto) > 255 or AVISO not in bruto.replace("\n", " "):
            rod_ruins.append((ws.title, len(bruto)))
        if not any(str(c.value or "").endswith(AVISO) for c in ws[3]):
            aviso_ruins.append(ws.title)
    check(g, f"{caminho.name}: rodapé com aviso completo e ≤255 caracteres em todas as abas", not rod_ruins, rod_ruins)
    check(g, f"{caminho.name}: aviso exato visível na linha 3 de todas as abas", not aviso_ruins, aviso_ruins)
    ext = [f for f in formulas if re.search(r"\[\d+\]|\.xlsx", f)]
    check(g, f"{caminho.name}: nenhuma fórmula aponta para outro arquivo", not ext, ext[:3])
    check(g, f"{caminho.name}: recálculo completo ao abrir (fullCalcOnLoad)", bool(wb.calculation.fullCalcOnLoad))
    imp = [ws.title for ws in wb.worksheets if ws.page_setup.paperSize not in (9, "9") or not ws.print_area]
    check(g, f"{caminho.name}: A4 e área de impressão em todas as abas", not imp, imp)
    return wb


def z_rodape(caminho, aba):
    wb = load_workbook(caminho, read_only=False)
    _ = wb
    with zipfile.ZipFile(caminho) as z:
        wbxml = z.read("xl/workbook.xml").decode()
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
        rid = re.search(r'<sheet[^>]*name="' + re.escape(aba.replace("&", "&amp;")) + r'"[^>]*r:id="(rId\d+)"', wbxml).group(1)
        alvo = re.search(r'Id="' + rid + r'"[^>]*Target="/?(?:xl/)?([^"]+)"', rels) or re.search(r'Target="/?(?:xl/)?([^"]+)"[^>]*Id="' + rid + '"', rels)
        xml = z.read("xl/" + alvo.group(1)).decode()
    m = re.search(r"<oddFooter>(.*?)</oddFooter>", xml, re.S)
    if not m:
        return ""
    import html
    return html.unescape(m.group(1))


# ------------------------------------------------------------------ cenários
def col(aba, cid):
    for i, c in enumerate(aba.colunas, start=1):
        if c.id == cid:
            return i
    raise KeyError(cid)


def escrever(ws, aba, linha, valores):
    for cid, v in valores.items():
        c = ws.cell(linha, col(aba, cid))
        c.value = v
        if isinstance(v, str):
            c.data_type = "s"


def recalc(wb, nome):
    ent = BUILD / f"{nome}.xlsx"
    sai = BUILD / f"{nome}-recalc.xlsx"
    wb.save(ent)
    LO.recalcular(ent, sai)
    return load_workbook(sai, data_only=True)


def valor(wbv, aba, cid, linha):
    return wbv[aba.nome].cell(linha, col(aba, cid)).value


def cenario_criticidade():
    g = "4 · Criticidade (LibreOffice recalculado)"
    A = E.ATIVOS
    casos = [("a1", [1, 1, 1, 1, 1], "C", 5), ("a2", [1, 2, 2, 2, 2], "B", 9), ("a3", [2, 2, 2, 3, 3], "A", 12), ("a4", [3, 1, 1, 1, 1], "A", 7),
             ("a5", [2, 2, None, 2, 2], "Pendente", None), ("a6", [4, 1, 1, 1, 1], "Pendente", None), ("a7", [2, 2, 2, 2, 2], "B", 10),
             ("a8", [2.5, 2, 2, 2, 2], "Pendente", None)]
    notas = ["criticidade.nota_seguranca", "criticidade.nota_qualidade", "criticidade.nota_producao", "criticidade.nota_frequencia_falha", "criticidade.nota_custo_reparo"]

    def montar(lim_a, lim_b):
        wb = load_workbook(MODELO["02"])
        ws = wb[A.nome]
        for k, (i, ns, _, _) in enumerate(casos):
            escrever(ws, A, 7 + k, {"ativo.id": i, "ativo.tag": f"UTL-MOT-00{k + 1}", "ativo.descricao": "Ativo fictício", "ativo.tipo_id": "motor_eletrico",
                                    "ativo.setor": "UTL", "ativo.numero": k + 1})
            for n, v in zip(notas, ns):
                ws.cell(7 + k, col(A, n)).value = v
        wb["Critérios"]["B7"].value = lim_a
        wb["Critérios"]["B8"].value = lim_b
        return wb

    wbv = recalc(montar(12, 9), "crit-padrao")
    for k, (i, ns, esperado, soma) in enumerate(casos):
        cl = valor(wbv, A, "criticidade.classe", 7 + k)
        pt = valor(wbv, A, "criticidade.pontuacao", 7 + k)
        check(g, f"limites 12/9: notas {ns} → {esperado}", cl == esperado and (pt == soma or (soma is None and pt in (None, ""))), f"classe={cl} pontuação={pt}")
    ori = valor(wbv, A, "criticidade.orientacao", 12)
    check(g, "nota 4 sinalizada como inválida na orientação", "inválida" in str(ori), ori)
    check(g, "linha vazia não mostra classe", valor(wbv, A, "criticidade.classe", 20) in (None, ""))
    wbv = recalc(montar(12, 12), "crit-limite-invalido")
    cls = [valor(wbv, A, "criticidade.classe", 7 + k) for k in range(len(casos))]
    check(g, "limite B ≥ A (12/12) → todas Pendente e aviso em Critérios", all(c == "Pendente" for c in cls) and
          "Inválidos" in str(wbv["Critérios"]["C9"].value), (cls, wbv["Critérios"]["C9"].value))
    wbv = recalc(montar(10, 8), "crit-limites-10-8")
    esperado = {"a1": "C", "a2": "B", "a3": "A", "a4": "A", "a7": "A"}
    obtido = {i: valor(wbv, A, "criticidade.classe", 7 + k) for k, (i, *_r) in enumerate(casos) if i in esperado}
    check(g, "alterar limites válidos (A=10, B=8) recalcula as classes", obtido == esperado, obtido)
    # validação de dados recusa 4 ao digitar
    wb = load_workbook(MODELO["02"])
    dvs = [dv for dv in wb[A.nome].data_validations.dataValidation if "K7" in str(dv.sqref)]
    check(g, "validação de dados da nota: inteiro de 1 a 3 com mensagem de erro (recusa 4 ao digitar)",
          bool(dvs) and dvs[0].type == "whole" and dvs[0].formula1 == "1" and dvs[0].formula2 == "3" and dvs[0].showErrorMessage)


def base_ocorrencias():
    wb = load_workbook(MODELO["03"])
    P, O = E.PLANO, E.OCORRENCIAS
    planos = [("T1", "Quinzenal", "Quinzenal", "Não", 1.0, "Sim"), ("TD", "Diária", "Diária", "Não", 0.25, "Sim"),
              ("TS1", "Mensal", "Trimestral", "Sim", 0.5, "Sim"), ("TS2", "Mensal", "Semanal", "Sim", 0.5, "Sim"), ("TS3", "Mensal", "Mensal", "Sim", 0.5, "Não")]
    for k, (t, ref, vig, seg, dur, ativa) in enumerate(planos):
        escrever(wb[P.nome], P, 7 + k, {"tarefa.ativo_id": "eq1", "tarefa.id": t, "tarefa.tag": "UTL-CMP-001",
                                        "tarefa.atividade": "Conteúdo técnico A definir", "tarefa.origem": "A definir",
                                        "tarefa.versao": "0.0.0", "tarefa.frequencia_referencia": ref, "tarefa.seguranca": seg, "tarefa.frequencia_vigente": vig,
                                        "tarefa.especialidade": "Mecânica", "tarefa.duracao_horas": dur, "tarefa.ativa": ativa})
    datas = [("T1", dt.date(2026, 1, 5)), ("T1", dt.date(2026, 1, 19)), ("T1", dt.date(2025, 12, 22)), ("T1", dt.date(2027, 1, 4))] + \
            [("TD", dt.date(2026, 2, 2) + dt.timedelta(days=i)) for i in range(5)]
    for k, (t, d) in enumerate(datas):
        escrever(wb[O.nome], O, 7 + k, {"ocorrencia.ativo_id": "eq1", "ocorrencia.tarefa_id": t, "ocorrencia.fonte": "Tabela de teste"})
        wb[O.nome].cell(7 + k, col(O, "ocorrencia.data_prevista")).value = d
    wb["Cronograma"]["C7"].value = 2026
    return wb


def execucao(wb, linha, id_, tarefa, prevista, real, cancelada="Não", origem="app", ativo="eq1", obs=""):
    X = E.EXECUCOES
    ws = wb[X.nome]
    escrever(ws, X, linha, {"execucao.id": id_, "execucao.ativo_id": ativo, "execucao.tarefa_id": tarefa, "execucao.tag": "UTL-CMP-001",
                            "execucao.descricao": "Descrição preservada fictícia", "execucao.origem": origem, "execucao.cancelada": cancelada,
                            "execucao.observacao": obs})
    ws.cell(linha, col(X, "execucao.data_prevista")).value = prevista
    ws.cell(linha, col(X, "execucao.data")).value = real


def resumo(wbv):
    ws = wbv["Cronograma"]
    j = E.CRONO["jan"]
    meses = {m + 1: (ws.cell(j + m, 3).value, ws.cell(j + m, 4).value, ws.cell(j + m, 5).value, ws.cell(j + m, 6).value) for m in range(12)}
    fora = {ws.cell(E.CRONO["fora_cab"] + k, 2).value: ws.cell(E.CRONO["fora_cab"] + k, 3).value for k in (1, 2)}
    semv = [ws.cell(E.CRONO["semv_cab"] + k, 3).value for k in (1, 2, 3)]
    return meses, fora, semv


def cenario_ocorrencias():
    g = "5 · Ocorrências e percentual (LibreOffice recalculado)"
    jan5, jan19 = dt.date(2026, 1, 5), dt.date(2026, 1, 19)
    wb = base_ocorrencias()
    execucao(wb, 7, "x1", "T1", jan5, jan5)
    m, fora, semv = resumo(recalc(wb, "oc-1"))
    check(g, "duas previstas e uma concluída → 50%", m[1][:3] == (2, 1, 0.5), m[1])
    check(g, "mês sem previsão (março) mostra 'Sem previsão'", m[3][2] == "Sem previsão" and m[3][0] == 0, m[3])
    wb = base_ocorrencias()
    execucao(wb, 7, "x1", "T1", jan5, jan5)
    execucao(wb, 8, "x2", "T1", jan5, dt.date(2026, 1, 6))
    m, _, _ = resumo(recalc(wb, "oc-duplicada"))
    check(g, "duplicar a execução da ocorrência concluída continua 50%", m[1][:3] == (2, 1, 0.5), m[1])
    wb = base_ocorrencias()
    execucao(wb, 7, "x1", "T1", jan5, jan5, cancelada="Sim")
    execucao(wb, 8, "x2", "T1", jan5, dt.date(2026, 1, 6), cancelada="Sim")
    wbv = recalc(wb, "oc-canceladas")
    m, _, semv = resumo(wbv)
    check(g, "cancelar todas as execuções vinculadas → 0%", m[1][:3] == (2, 0, 0), m[1])
    check(g, "canceladas continuam no histórico e aparecem no quadro separado", semv[2] == 2 and wbv[E.EXECUCOES.nome]["A7"].value == "x1", semv)
    wb = base_ocorrencias()
    execucao(wb, 7, "x1", "T1", jan5, jan5)
    execucao(wb, 8, "x3", "T1", None, dt.date(2026, 1, 10), origem="legado")
    m, _, semv = resumo(recalc(wb, "oc-legado"))
    check(g, "legado sem vínculo não aumenta o percentual e aparece separado", m[1][:3] == (2, 1, 0.5) and semv[0] == 1, (m[1], semv))
    wb = base_ocorrencias()
    execucao(wb, 7, "x4", "T1", jan19, dt.date(2026, 2, 3))
    m, _, _ = resumo(recalc(wb, "oc-outro-mes"))
    check(g, "executar em outro mês conta no mês previsto (jan 1/2, fev 0/5)", m[1][:2] == (2, 1) and m[2][:2] == (5, 0), (m[1], m[2]))
    g6 = "6 · Ano, diária, histórico e vazios"
    wb = base_ocorrencias()
    execucao(wb, 7, "x5", "T9", None, dt.date(2026, 3, 2), ativo="eq-removido")
    wbv = recalc(wb, "oc-ano-diaria")
    F = E.FALHAS
    wbf = load_workbook(MODELO["05"])
    escrever(wbf[F.nome], F, 7, {"falha.id": "f1", "falha.ativo_id": "eq1", "falha.tag": "UTL-CMP-001", "falha.sintoma": "Sintoma fictício"})
    wbf[F.nome].cell(7, col(F, "falha.data")).value = dt.date(2026, 3, 1)
    escrever(wbf[F.nome], F, 8, {"falha.id": "f2", "falha.ativo_id": "eq1", "falha.tag": "UTL-CMP-001", "falha.sintoma": "Sintoma fictício", "falha.horas_parada": 0})
    wbf[F.nome].cell(8, col(F, "falha.data")).value = dt.date(2026, 3, 2)
    wbfv = recalc(wbf, "falhas-vazio-zero")
    m, fora, _ = resumo(wbv)
    check(g6, "passagem de ano: previsões de 2025 e 2027 ficam fora do ano (1 antes, 1 depois)", fora == {"Antes do ano": 1, "Depois do ano": 1}, fora)
    check(g6, "diária: 5 dias úteis = 5 ocorrências e 1,25 h (sem multiplicar por 5 de novo)", m[2][0] == 5 and abs(m[2][3] - 1.25) < 1e-9, m[2])
    X = E.EXECUCOES
    check(g6, "ativo/tarefa removidos do plano: histórico preservado com TAG e aviso",
          valor(wbv, X, "execucao.situacao_plano", 7) == "Fora do plano: histórico preservado"
          and valor(wbv, X, "execucao.tag", 7) == "UTL-CMP-001", valor(wbv, X, "execucao.situacao_plano", 7))
    p1, p2 = valor(wbfv, F, "falha.parada_situacao", 7), valor(wbfv, F, "falha.parada_situacao", 8)
    h1, h2 = valor(wbfv, F, "falha.horas_parada", 7), valor(wbfv, F, "falha.horas_parada", 8)
    check(g6, "horas de parada vazias ≠ zero (Não informado × Informado 0)", p1 == "Não informado" and h1 is None and p2 == "Informado" and h2 == 0, (p1, h1, p2, h2))
    meta = wbv["Cronograma"][E.CRONO["card_meta"]].value
    wbo = recalc(load_workbook(MODELO["01"]), "organizacao-vazia")
    check(g6, "ausência de meta mostra 'A definir' (nenhuma meta inventada)", meta == "A definir" and wbo["Organização"]["C12"].value == "A definir",
          (meta, wbo["Organização"]["C12"].value))
    g7 = "7 · Segurança e capacidade"
    P = E.PLANO
    confs = [valor(wbv, P, "tarefa.conferencia", 7 + k) for k in range(5)]
    check(g7, "segurança com intervalo ampliado (Mensal→Trimestral) → ERRO visível", str(confs[2]).startswith("ERRO: intervalo ampliado"), confs[2])
    check(g7, "segurança com frequência aumentada (Mensal→Semanal) → OK", str(confs[3]).startswith("OK"), confs[3])
    check(g7, "segurança retirada do plano (Ativa = Não) → ERRO visível", str(confs[4]).startswith("ERRO: tarefa de segurança retirada"), confs[4])
    check(g7, "referência técnica do plano protegida (aba protegida, colunas REFERÊNCIA bloqueadas, sem senha)",
          wb[P.nome].protection.sheet and wb[P.nome].cell(7, col(P, "tarefa.frequencia_referencia")).protection.locked
          and not wb[P.nome].cell(7, col(P, "tarefa.frequencia_vigente")).protection.locked and not wb[P.nome].protection.password)


def cenario_texto_capacidade():
    g3 = "3 · Tipos de dados"
    wb = base_ocorrencias()
    perigosos = ["=1+1", "+55 texto", "-texto", "@texto", "=SOMA(A1:A2)"]
    for k, t in enumerate(perigosos):
        execucao(wb, 7 + k, f"t{k}", "T1", None, dt.date(2026, 1, 2), obs=t)
    A = E.ATIVOS
    wb[E.OCORRENCIAS.nome].cell(7 + E.OCORRENCIAS.capacidade, 1).value = "fora"
    wbv = recalc(wb, "texto-capacidade")
    wba = load_workbook(MODELO["02"])
    wba[A.nome].cell(7 + A.capacidade, 1).value = "fora"
    wbav = recalc(wba, "capacidade-ativos")
    X = E.EXECUCOES
    obtidos = [valor(wbv, X, "execucao.observacao", 7 + k) for k in range(len(perigosos))]
    check(g3, "textos iniciados por =, +, - e @ continuam texto após recálculo no LibreOffice", obtidos == perigosos, obtidos)
    g7 = "7 · Segurança e capacidade"
    linha = next((r for r in reversed(list(wbav["Config"].iter_rows())) if r[0].value == A.nome), None)  # a última é a das Conferências
    alerta = linha[3].value if linha else None
    check(g7, "dado abaixo da capacidade de Ativos gera alerta ATENÇÃO nas Conferências da aba Config (não falha em silêncio)",
          str(alerta).startswith("ATENÇÃO"), alerta)
    c = wbv["Cronograma"].cell(E.CRONO["conf_cab"] + 3, 3).value
    check(g7, "dado abaixo da capacidade de Ocorrências gera alerta no PCMD-04 Cronograma", str(c).startswith("ATENÇÃO"), c)


def verificar_exemplo(demo):
    g3 = "3 · Tipos de dados"
    X, F, P, A, O = E.EXECUCOES, E.FALHAS, E.PLANO, E.ATIVOS, E.OCORRENCIAS
    wb3, wb5, wb2 = load_workbook(EXEMPLO["03"]), load_workbook(EXEMPLO["05"]), load_workbook(EXEMPLO["02"])
    ws = wb3[X.nome]
    datas_ok = all(isinstance(ws.cell(r, col(X, "execucao.data")).value, dt.datetime) and ws.cell(r, col(X, "execucao.data")).number_format == "dd/mm/yyyy"
                   for r in range(7, 7 + len(demo["gestao"]["historico"])))
    oc_ok = all(isinstance(wb3[O.nome].cell(r, 3).value, dt.datetime) for r in range(7, 7 + len(demo["ocorrencias"])))
    check(g3, "datas são datas reais do Excel com formato dd/mm/aaaa (Execuções e Ocorrências)", datas_ok and oc_ok)
    horas = [wb3[P.nome].cell(r, col(P, "tarefa.duracao_horas")).value for r in range(7, 7 + len(demo["plano"]))]
    check(g3, "horas são números decimais (Plano Mestre)", all(isinstance(h, (int, float)) for h in horas))
    fh = [wb5[F.nome].cell(r, col(F, "falha.horas_parada")).value for r in range(7, 10)]
    check(g3, "horas de parada: 1,5 / vazio / 0 preservados (vazio não virou zero)", fh == [1.5, None, 0], fh)
    with zipfile.ZipFile(EXEMPLO["03"]) as z:
        ss = z.read("xl/sharedStrings.xml").decode() if "xl/sharedStrings.xml" in z.namelist() else ""
        sh = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
    perig = [h["observacao"] for h in demo["gestao"]["historico"] if h["observacao"] and h["observacao"][0] in "=+-@"]
    import html
    texto_xml = html.unescape(ss + sh)  # o openpyxl grava textos inline com entidades (&#227;)
    presentes = all(f"<t>{p}</t>" in texto_xml for p in perig)
    sem_formula = all(f"<f>{p[1:]}</f>" not in texto_xml for p in perig)
    check(g3, "textos perigosos do exemplo gravados como texto, não como fórmula", presentes and sem_formula, perig)
    ids_xlsx = [wb3[X.nome].cell(r, 1).value for r in range(7, 7 + len(demo["gestao"]["historico"]))]
    ids_eq = [wb2[A.nome].cell(r, 1).value for r in range(7, 7 + len(demo["equipamentos"]))]
    check(g3, "IDs técnicos intactos (execuções e ativos iguais aos dados de origem)",
          ids_xlsx == [h["id"] for h in demo["gestao"]["historico"]] and ids_eq == [e["id"] for e in demo["equipamentos"]])
    g8 = "8 · Recálculo do exemplo comparado com a lógica do app"
    recalc_ex = {}
    for k in ARQ:
        sai = BUILD / f"exemplo-{k}-recalc.xlsx"
        LO.recalcular(EXEMPLO[k], sai)
        recalc_ex[k] = load_workbook(sai, data_only=True)
    wbv = recalc_ex["03"]
    m, fora, semv = resumo(wbv)
    app = demo["resumoMensal"]
    dif = []
    for mes in app["meses"]:
        p, r, pc, h = m[mes["mes"]]
        esp_pc = "Sem previsão" if mes["percentual"] is None else mes["percentual"] / 100
        if p != mes["previstas"] or r != mes["realizadas"] or abs((h or 0) - mes["horas"]) > 1e-9 or \
                (esp_pc == "Sem previsão" and pc != esp_pc) or (esp_pc != "Sem previsão" and abs(pc - esp_pc) > 1e-9):
            dif.append((mes, m[mes["mes"]]))
    check(g8, "12 meses: previstas, concluídas, % e horas iguais a resumoMensal() do app", not dif, dif[:2])
    check(g8, "fora do ano igual ao app (previstas)", (fora["Antes do ano"] or 0) + (fora["Depois do ano"] or 0) == app["foraDoAno"]["previstas"], fora)
    check(g8, "sem vínculo no ano igual ao app", semv[0] == app["semVinculo"], semv)
    tot = sum(x["previstas"] for x in app["meses"])
    total = wbv["Cronograma"].cell(E.CRONO["total"], 3).value
    check(g8, "Total do ano (previstas) igual à soma dos meses do app", total == tot, (total, tot))
    w2 = recalc_ex["02"]
    cls = {w2[A.nome].cell(r, 1).value: (w2[A.nome].cell(r, col(A, "criticidade.pontuacao")).value, w2[A.nome].cell(r, col(A, "criticidade.classe")).value)
           for r in range(7, 7 + len(demo["equipamentos"]))}
    esp = {k: (v["pontuacao"], v["classe"]) for k, v in demo["criticidade"].items()}
    cls = {k: (None if v[0] in ("", None) else v[0], v[1]) for k, v in cls.items()}
    check(g8, "criticidade do exemplo igual a criticidadeAtivo() do app", cls == esp, (cls, esp))
    erros = []
    for k, arq in ARQ.items():
        for aba in arq.tabelas:
            wsv = recalc_ex[k][aba.nome]
            for r in range(7, wsv.max_row + 1):
                for c in range(1, len(aba.colunas) + 1):
                    v = wsv.cell(r, c).value
                    if isinstance(v, str) and (v.startswith("#") or v.startswith("Err:")):
                        erros.append((aba.nome, r, c, v))
        for ws_ in recalc_ex[k].worksheets:  # topo, formulários e Cronograma
            for row in ws_.iter_rows(min_row=1, max_row=min(ws_.max_row, 80)):
                for c in row:
                    if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:")):
                        erros.append((ws_.title, c.coordinate, c.value))
    check(g8, "nenhum erro de fórmula (#VALOR!, #N/D, Err:) nos quatro exemplos recalculados", not erros, erros[:5])
    confs = [wbv[P.nome].cell(r, col(P, "tarefa.conferencia")).value for r in range(7, 7 + len(demo["plano"]))]
    check(g8, "plano do exemplo sem ERRO de segurança", not any(str(c).startswith("ERRO") for c in confs))


# ------------------------------------------------------------------ 10: estrutura igual e contrato
def assinatura(caminho):
    wb = load_workbook(caminho)
    s = {"abas": wb.sheetnames, "nomes": sorted((n, d.attr_text) for n, d in wb.defined_names.items()), "cab": {}, "formulas": {}, "dv": {}}
    for ws in wb.worksheets:
        s["cab"][ws.title] = [[ws.cell(r, c).value for c in range(1, 25)] for r in (5, 6)]
        s["dv"][ws.title] = sorted((dv.type, dv.formula1, str(dv.sqref)) for dv in ws.data_validations.dataValidation)
        f = {}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and c.data_type == "f":
                    f[c.coordinate] = c.value
        s["formulas"][ws.title] = f
    return s


def verificar_contrato():
    g = "10 · Contrato e estrutura"
    for k, arq in ARQ.items():
        a, b = assinatura(MODELO[k]), assinatura(EXEMPLO[k])
        iguais = [ch for ch in ("abas", "nomes", "cab", "dv") if a[ch] == b[ch]]
        check(g, f"{arq.nome}: modelo e exemplo com a mesma estrutura (abas, nomes, cabeçalhos, validações)", len(iguais) == 4, iguais)
        dif = [(aba, kk) for aba in a["formulas"] for kk in a["formulas"][aba] if b["formulas"][aba].get(kk) != a["formulas"][aba][kk]]
        total = sum(len(v) for v in a["formulas"].values())
        check(g, f"{arq.nome}: mesmas fórmulas nas mesmas células ({total})", not dif, dif[:3])
    contrato = json.loads((RAIZ / "CONTRATO-CAMPOS.json").read_text(encoding="utf-8"))
    check(g, "CONTRATO-CAMPOS.json é JSON válido, versão 1", contrato.get("versao") == 1)
    xlsx = [d for d in contrato["arquivos"] if d["arquivo"].endswith(".xlsx")]
    check(g, "contrato descreve os quatro arquivos XLSX", sorted(d["arquivo"] for d in xlsx) == sorted(f"planilhas/{a.nome}.xlsx" for a in ARQ.values()))
    faltando, nomes_ruins, sem = [], [], []
    campos_obrig = ["id", "rotulo", "tipo", "obrigatorio", "formato_unidade", "origem_proposta", "destino", "regra_vazio", "validacao",
                    "entrada_ou_calculo", "privacidade", "observacoes"]
    for d in xlsx:
        wb = load_workbook(RAIZ / d["arquivo"])
        for aba in d["abas_tabela"]:
            ws = wb[aba["nome"]]
            for c in aba["colunas"]:
                letra_ = c["destino"]["coluna"]
                if ws[f"{letra_}{aba['linha_cabecalho']}"].value != c["rotulo"]:
                    faltando.append((d["arquivo"], aba["nome"], letra_))
                if any(kk not in c for kk in campos_obrig):
                    sem.append((aba["nome"], c["id"]))
            ncab = sum(1 for cc in range(1, 40) if ws.cell(aba["linha_cabecalho"], cc).value)
            if ncab != len(aba["colunas"]):
                faltando.append((aba["nome"], "nº de colunas", ncab, len(aba["colunas"])))
        for c in d["celulas_unicas"]:
            if c["destino"]["aba"] not in wb.sheetnames:
                faltando.append((d["arquivo"], c["id"]))
        if {n["nome"] for n in d["intervalos_nomeados"]} != set(wb.defined_names.keys()):
            nomes_ruins.append(d["arquivo"])
    check(g, "contrato cobre todas as colunas de todas as tabelas XLSX (rótulo e posição conferidos)", not faltando, faltando[:5])
    check(g, "contrato lista todos os intervalos nomeados de cada arquivo", not nomes_ruins, nomes_ruins)
    check(g, "cada campo do contrato tem identificador, rótulo, tipo, obrigatoriedade, formato, origem, destino, vazio, validação, papel e privacidade", not sem, sem[:3])
    for d in [x for x in contrato["arquivos"] if x["arquivo"].endswith(".docx")]:
        caminho = RAIZ / d["arquivo"]
        with zipfile.ZipFile(caminho) as z:
            xml = "".join(z.read(n).decode() for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml", n))
        encontrados = set(re.findall(r"\{\{[a-z0-9_.]+\}\}", xml))
        listados = {m["marcador"] for m in d["marcadores"]} | {x for r in d["regioes_repetiveis"] for x in r["marcadores_da_linha"]}
        check(g, f"{caminho.name}: marcadores do arquivo = marcadores do contrato ({len(encontrados)})", encontrados == listados,
              (encontrados ^ listados))
        quebrados = re.findall(r"\{\{[^}<]*<", xml)
        check(g, f"{caminho.name}: nenhum marcador dividido entre trechos de texto", not quebrados, quebrados[:2])
        locais = {l for m in d["marcadores"] for l in m["locais"]}
        check(g, f"{caminho.name}: marcadores em parágrafo, célula, cabeçalho e rodapé documentados",
              {"parágrafo", "cabeçalho", "rodapé"} <= locais and any("célula" in l for l in locais), sorted(locais))


# ------------------------------------------------------------------ DOCX
def verificar_docx():
    g = "1/9 · DOCX"
    for nome, caminho in DOCX.items():
        nomes, xml = estrutura_zip(caminho, "2 · Estrutura DOCX")
        doc = Document(caminho)
        check(g, f"{caminho.name}: abre no python-docx", doc is not None)
        campos = set(re.findall(r'w:instr="\s*([A-Z]+)', xml)) | set(re.findall(r"<w:instrText[^>]*>\s*([A-Z]+)", xml))
        check(g, f"{caminho.name}: apenas campos PAGE e NUMPAGES (sem LINK/INCLUDE)", campos <= {"PAGE", "NUMPAGES"}, campos)
        estilos = {p.style.name for p in doc.paragraphs}
        check(g, f"{caminho.name}: usa estilos reais de título (Heading 1, Heading 2)", {"Heading 1", "Heading 2"} <= estilos)
        secs = doc.sections
        orient = {str(s.orientation) for s in secs}
        rodape_ok = AVISO in "\n".join(p.text for p in secs[0].footer.paragraphs) and all(s.footer.is_linked_to_previous for s in secs[1:])
        check(g, f"{caminho.name}: aviso no rodapé da 1ª seção e demais seções vinculadas ({len(secs)} seções, orientações {sorted(orient)})",
              rodape_ok and len(orient) == 2)
        tbl_header = xml.count("<w:tblHeader")
        check(g, f"{caminho.name}: tabelas com linha de cabeçalho repetida ({tbl_header} de {len(doc.tables)}; a grade de identificação não tem cabeçalho)",
              tbl_header >= len(doc.tables) - 1 and len(doc.tables) > 5)
        if "Exemplo" in nome:
            check(g, f"{caminho.name}: nenhum marcador restante no exemplo", "{{" not in xml)
            check(g, f"{caminho.name}: aprovação não declarada", "aprovação não registrada" in xml)
    tem_logo = all("word/media/" in "".join(zipfile.ZipFile(c).namelist()) for c in DOCX.values())
    check(g, "logo do modelo padrão PCM Direto no cabeçalho dos quatro DOCX", tem_logo)


# ------------------------------------------------------------------ PDF (9)
REPETIDAS = {"03": (("Plano Mestre", "ID do ativo"), ("Ocorrências", "ID do ativo"), ("Execuções", "ID do registro")),
             "02": (("Ativos", "ID do ativo"),), "05": (("Ficha de Falha", "ID do registro"),)}


def verificar_pdfs():
    g = "9 · Impressão (PDF do LibreOffice)"
    alvos = [(k, c) for k in ARQ for c in (MODELO[k], EXEMPLO[k])] + [(None, c) for c in DOCX.values()]
    for k, caminho in alvos:
        pdf = PREVIAS / (caminho.stem + ".pdf")
        LO.pdf(caminho, pdf)
        n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1))
        sem_aviso = []
        paginas = []
        for p in range(1, n + 1):
            t = subprocess.run(["pdftotext", "-layout", "-f", str(p), "-l", str(p), str(pdf), "-"], capture_output=True, text=True).stdout
            paginas.append(t)
            if norm(AVISO) not in norm(t):
                sem_aviso.append(p)
        check(g, f"{pdf.name}: aviso completo em todas as {n} páginas", not sem_aviso, sem_aviso[:10])
        if caminho.suffix == ".xlsx":
            for aba, marca in REPETIDAS.get(k, ()):
                pags = [t for t in paginas if f"· {aba}" in t.split("\n", 2)[0] or re.search(rf"PCMD-0\d · {aba}\b", t[:300])]
                faltam = [i for i, t in enumerate(pags) if marca not in t]
                if pags:
                    check(g, f"{pdf.name}: cabeçalho da tabela repetido em todas as {len(pags)} páginas de {aba}", not faltam, faltam[:5])
        else:
            marcas = "{{" in "".join(paginas)
            if "Exemplo" not in caminho.name:
                check(g, f"{pdf.name}: marcadores visíveis no corpo, no cabeçalho e no rodapé impressos",
                      marcas and "{{organizacao.nome}}" in paginas[-1] and "{{documento.revisao}}" in paginas[-1])


def main():
    demo = json.loads((RAIZ / "build" / "demo.json").read_text(encoding="utf-8"))
    print("LibreOffice:", LO.versao())
    g1 = "1 · Abertura"
    for c in list(MODELO.values()) + list(EXEMPLO.values()) + list(DOCX.values()):
        try:
            d = LO.abrir(c)
            d.close(True)
            check(g1, f"{c.name}: aberto pelo LibreOffice headless (RepairPackage desligado)", True)
        except Exception as e:  # noqa: BLE001
            check(g1, f"{c.name}: aberto pelo LibreOffice headless", False, e)
    for k in ARQ:
        verificar_xlsx_estrutura(MODELO[k])
        verificar_xlsx_estrutura(EXEMPLO[k])
    verificar_docx()
    verificar_exemplo(demo)
    cenario_criticidade()
    cenario_ocorrencias()
    cenario_texto_capacidade()
    verificar_contrato()
    verificar_pdfs()
    LO.encerrar()
    saida = {"executado_em": dt.datetime.now().isoformat(timespec="seconds"), "libreoffice": LO.versao(),
             "total": len(RESULTADOS), "falhas": sum(r["resultado"] == "FALHOU" for r in RESULTADOS), "resultados": RESULTADOS}
    (AQUI / "resultado-testes.json").write_text(json.dumps(saida, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"\nTotal: {saida['total']} · falhas: {saida['falhas']}")
    return 1 if saida["falhas"] else 0


if __name__ == "__main__":
    sys.exit(main())
