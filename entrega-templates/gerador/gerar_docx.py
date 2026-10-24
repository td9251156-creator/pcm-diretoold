"""Gera os manuais PCMD-W1 e PCMD-W2 (templates com marcadores e versões de exemplo).

Template e exemplo saem da mesma função: no template cada {{marcador}} fica visível, em um
único trecho de texto (run), para facilitar um preenchimento futuro; no exemplo o marcador é
substituído por um valor fictício e as regiões repetíveis recebem uma linha por registro.
"""
import datetime as dt
import json
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import especificacao as E
from comum import (AUTOR, AVISO, AZUL, CAB_FUNDO, CINZA_APOIO, CONTATO, DESTAQUE, FONTE, FONTE_ESTREITA, FONTE_TITULO, LICENCA, LOGO_MARCA,
                   RAIZ_REPO, SECAO_BORDA, SECAO_FUNDO, SEGURANCA, TEXTO, VERSAO_MODELO)

BORDA_TABELA = "D5DCE8"

MARCADOR = re.compile(r"(\{\{[a-z0-9_.]+\}\})")


def _cor(hexa):
    return RGBColor.from_string(hexa)


def sombrear(celula, cor):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def campo(paragrafo, instrucao, texto_padrao="1"):
    """Campo simples do Word (PAGE, NUMPAGES)."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instrucao)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto_padrao
    r.append(t)
    fld.append(r)
    paragrafo._p.append(fld)


def borda_paragrafo(pPr, lado, cor, tamanho=8, espaco=4):
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(tamanho))
    el.set(qn("w:space"), str(espaco))
    el.set(qn("w:color"), cor)
    bdr.append(el)


def sombra_paragrafo(pPr, cor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    pPr.append(shd)


def bordas_tabela(tabela, cor=BORDA_TABELA, tamanho=4):
    """Grade fina em todas as linhas, como no modelo padrão de documento."""
    tblPr = tabela._tbl.tblPr
    antigo = tblPr.find(qn("w:tblBorders"))
    if antigo is not None:
        tblPr.remove(antigo)
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(tamanho))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), cor)
        b.append(el)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("top", 50), ("bottom", 50), ("left", 90), ("right", 90)):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def borda_celula(celula, lado, cor, tamanho):
    tcPr = celula._tc.get_or_add_tcPr()
    bd = tcPr.find(qn("w:tcBorders"))
    if bd is None:
        bd = OxmlElement("w:tcBorders")
        tcPr.append(bd)
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(tamanho))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), cor)
    bd.append(el)


def fonte_run(run, nome):
    rpr = run._r.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts")
        rpr.insert(0, f)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        f.set(qn(a), nome)


def tabela_fixa(t, larguras_cm):
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(larguras_cm):
        t.columns[i].width = Cm(w)  # grade da tabela (tblGrid): respeitada por Word e LibreOffice
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)


def repetir_cabecalho(linha):
    trPr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def nao_quebrar(linha):
    trPr = linha._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


class Manual:
    def __init__(self, codigo, titulo, valores=None, regioes=None):
        self.codigo, self.titulo = codigo, titulo
        self.valores = valores  # None = template
        self.regioes = regioes or {}
        self.doc = Document()
        self.marcadores_usados = []  # (marcador, local)
        self.regioes_usadas = []
        self._estilos()
        self._pagina(self.doc.sections[0])
        self._cabecalho_rodape(self.doc.sections[0])

    # ------------------------------------------------------------ estilos e página
    def _estilos(self):
        """Estilos do Modelo Padrão de Documento PCM Direto."""
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = FONTE
        normal.font.size = Pt(10)
        normal.font.color.rgb = _cor(TEXTO)
        rpr = normal.element.get_or_add_rPr()
        fontes = rpr.find(qn("w:rFonts"))
        if fontes is None:
            fontes = OxmlElement("w:rFonts")
            rpr.append(fontes)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fontes.set(qn(a), FONTE)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "pt-BR")
        rpr.append(lang)
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for nome, tam, fam, negrito in (("Heading 1", 11, FONTE_TITULO, False), ("Heading 2", 10.5, FONTE, True), ("Heading 3", 10, FONTE, True)):
            s = st[nome]
            s.font.name = fam
            s.font.size = Pt(tam)
            s.font.bold = negrito
            s.font.italic = False
            s.font.color.rgb = _cor(AZUL)
            s_rpr = s.element.get_or_add_rPr()
            f = s_rpr.find(qn("w:rFonts"))
            if f is None:
                f = OxmlElement("w:rFonts")
                s_rpr.append(f)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                f.set(qn(a), fam)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if f.get(qn(a)) is not None:
                    del f.attrib[qn(a)]
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            s.paragraph_format.space_before = Pt(12 if nome == "Heading 1" else 8)
            s.paragraph_format.space_after = Pt(5 if nome == "Heading 1" else 3)
        borda_paragrafo(st["Heading 1"].element.get_or_add_pPr(), "bottom", AZUL, 6, 2)
        for nome in ("List Bullet", "List Number"):
            st[nome].font.name = FONTE
        if "Marcador PCMD" not in [s.name for s in st]:
            m = st.add_style("Marcador PCMD", WD_STYLE_TYPE.CHARACTER)
            m.font.color.rgb = _cor(DESTAQUE)
            m.font.bold = True
        av = st.add_style("Aviso PCMD", WD_STYLE_TYPE.PARAGRAPH)
        av.base_style = st["Normal"]
        av.font.size = Pt(8)
        av.font.name = FONTE_ESTREITA
        av.font.color.rgb = _cor(CINZA_APOIO)
        av.paragraph_format.space_after = Pt(0)
        av.paragraph_format.line_spacing = 1.0
        av.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        nt = st.add_style("Nota PCMD", WD_STYLE_TYPE.PARAGRAPH)
        nt.base_style = st["Normal"]
        nt.font.size = Pt(9)
        nt.font.color.rgb = _cor(CINZA_APOIO)
        nt.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        de = st.add_style("Destaque PCMD", WD_STYLE_TYPE.PARAGRAPH)
        de.base_style = st["Normal"]
        de.font.size = Pt(9.5)
        de.font.name = FONTE_ESTREITA
        de.paragraph_format.space_before = Pt(4)
        de.paragraph_format.space_after = Pt(8)
        de.paragraph_format.left_indent = Cm(0.25)
        de.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tb = st.add_style("Tabela PCMD", WD_STYLE_TYPE.PARAGRAPH)
        tb.base_style = st["Normal"]
        tb.font.size = Pt(9)
        tb.paragraph_format.space_after = Pt(0)
        tb.paragraph_format.line_spacing = 1.0
        tb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _pagina(self, sec, paisagem=False):
        a4 = (Cm(21.0), Cm(29.7))
        if paisagem:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = a4[1], a4[0]
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = a4
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(2.29)
        sec.header_distance = Cm(0.5)
        sec.footer_distance = Cm(1.2)

    def destaque(self, texto):
        """Caixa NOTA do modelo padrão: barra azul à esquerda, fundo claro e rótulo NOTA."""
        par = self.doc.add_paragraph(style="Destaque PCMD")
        pPr = par._p.get_or_add_pPr()
        borda_paragrafo(pPr, "left", AZUL, 24, 6)
        sombra_paragrafo(pPr, SECAO_FUNDO)
        r = par.add_run("NOTA   ")
        r.font.size, r.font.color.rgb = Pt(8), _cor(AZUL)
        fonte_run(r, FONTE_TITULO)
        self._texto(par, texto, "parágrafo")
        return par

    def _cabecalho_rodape(self, sec):
        """Cabeçalho: logo | título | revisão e elaborado/aprovado; faixa azul abaixo. Rodapé: aviso e paginação."""
        largura = sec.page_width - sec.left_margin - sec.right_margin
        cab = sec.header
        larg = [4.85, 7.9, 4.25]  # 2750/4475/2413 dxa do modelo
        t = cab.add_table(rows=1, cols=3, width=Cm(sum(larg)))
        tabela_fixa(t, larg)
        bordas_tabela(t, SECAO_BORDA, 6)
        c_logo, c_tit, c_rev = t.rows[0].cells
        for c, w in zip(t.rows[0].cells, larg):
            c.width = Cm(w)
            c.vertical_alignment = 1  # centro
        p = c_logo.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if LOGO_MARCA.exists():
            p.add_run().add_picture(str(LOGO_MARCA), width=Cm(4.2))  # só a largura: preserva a proporção
        else:
            p.add_run("PCM Direto").bold = True
        p = c_tit.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(self.titulo.upper())
        r.font.size, r.font.color.rgb = Pt(12), _cor("000000")
        fonte_run(r, FONTE_TITULO)
        p2 = c_tit.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(self.codigo)
        r.font.size, r.bold, r.font.color.rgb = Pt(8.5), True, _cor(DESTAQUE)
        p = c_rev.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n0 = len(p.runs)
        r = p.add_run("Rev. ")
        self._texto(p, "{{documento.revisao}}", "cabeçalho")
        p.add_run(" · ")
        self._texto(p, "{{documento.data}}", "cabeçalho")
        for r in p.runs[n0:]:
            r.font.size = Pt(8)
            fonte_run(r, FONTE_TITULO)
        p = c_rev.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Elaborado: ")
        self._texto(p, "{{documento.elaborador}}", "cabeçalho")
        p.add_run(" · Aprovado: ")
        self._texto(p, "{{documento.aprovador}}", "cabeçalho")
        for r in p.runs:
            r.font.size, r.bold = Pt(8), True
            if r.style.name != "Marcador PCMD":
                r.font.color.rgb = _cor(CINZA_APOIO)
            fonte_run(r, FONTE_ESTREITA)
        for c in t.rows[0].cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
        faixa = cab.paragraphs[0]  # parágrafo inicial do cabeçalho: vai para depois da tabela
        t._tbl.addnext(faixa._p)
        faixa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = faixa.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(17)
        sombra_paragrafo(faixa._p.get_or_add_pPr(), AZUL)
        n0 = len(faixa.runs)
        faixa.add_run("MANUAL · ")
        self._texto(faixa, "{{organizacao.nome}}", "cabeçalho")
        for r in faixa.runs[n0:]:
            r.font.size = Pt(9)
            if r.style.name != "Marcador PCMD":
                r.font.color.rgb = _cor("FFFFFF")
            fonte_run(r, FONTE_TITULO)
        rod = sec.footer
        p = rod.paragraphs[0]
        p.style = self.doc.styles["Aviso PCMD"]
        borda_paragrafo(p._p.get_or_add_pPr(), "top", SECAO_BORDA, 6, 4)
        p.add_run(AVISO)
        p = rod.add_paragraph(style="Aviso PCMD")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(largura, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(f"PCM Direto · {self.codigo} · Rev. ").font.size = Pt(8.5)
        self._texto(p, "{{documento.revisao}}", "rodapé", tamanho=8.5)
        p.add_run(f" · Licença {LICENCA}\tPág. ").font.size = Pt(8.5)
        campo(p, "PAGE")
        p.add_run(" de ").font.size = Pt(8.5)
        campo(p, "NUMPAGES")

    # ------------------------------------------------------------ texto com marcadores
    def valor(self, marcador):
        chave = marcador.strip("{}")
        v = self.valores.get(chave)
        return "A definir" if v in (None, "") else str(v)

    def _texto(self, paragrafo, texto, local, tamanho=None, negrito=False):
        for parte in MARCADOR.split(texto):
            if not parte:
                continue
            if MARCADOR.fullmatch(parte):
                self.marcadores_usados.append((parte.strip("{}"), local))
                if self.valores is None:
                    r = paragrafo.add_run(parte, style="Marcador PCMD")
                else:
                    r = paragrafo.add_run(self.valor(parte))
            else:
                r = paragrafo.add_run(parte)
            if tamanho:
                r.font.size = Pt(tamanho)
            if negrito:
                r.bold = True
        return paragrafo

    def p(self, texto, estilo=None):
        return self._texto(self.doc.add_paragraph(style=estilo), texto, "parágrafo")

    def h1(self, texto, numerar=True):
        """Seção do modelo padrão: «1.  OBJETIVO»."""
        if numerar:
            self._n_secao = getattr(self, "_n_secao", 0) + 1
            self._n_sub = 0
            texto = f"{self._n_secao}.  {texto}"
        return self.doc.add_heading(texto.upper(), level=1)

    def h2(self, texto):
        """Subseção numerada: «5.1  Preparação»."""
        self._n_sub = getattr(self, "_n_sub", 0) + 1
        return self.doc.add_heading(f"{self._n_secao}.{self._n_sub}  {texto}", level=2)

    def lista(self, itens, numerada=False):
        for i in itens:
            self.p(i, "List Number" if numerada else "List Bullet")

    def nota(self, texto):
        return self.p(texto, "Nota PCMD")

    # ------------------------------------------------------------ tabelas
    def _tabela(self, cabecalhos, larguras_cm):
        t = self.doc.add_table(rows=1, cols=len(cabecalhos))
        t.style = self.doc.styles["Table Grid"]
        tabela_fixa(t, larguras_cm)
        bordas_tabela(t)
        linha = t.rows[0]
        repetir_cabecalho(linha)
        nao_quebrar(linha)
        for i, c in enumerate(cabecalhos):
            cel = linha.cells[i]
            cel.width = Cm(larguras_cm[i])
            sombrear(cel, CAB_FUNDO)
            borda_celula(cel, "bottom", AZUL, 8)
            par = cel.paragraphs[0]
            par.style = self.doc.styles["Tabela PCMD"]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(c)
            r.bold = True
            r.font.color.rgb = _cor(AZUL)
        return t

    def _linha(self, t, textos, larguras_cm, local="célula de tabela", destaque=None, primeira_negrito=False):
        linha = t.add_row()
        nao_quebrar(linha)
        for i, tx in enumerate(textos):
            cel = linha.cells[i]
            cel.width = Cm(larguras_cm[i])
            par = cel.paragraphs[0]
            par.style = self.doc.styles["Tabela PCMD"]
            self._texto(par, tx, local)
            if primeira_negrito and i == 0:
                for r in par.runs:
                    r.bold = True
            if destaque and destaque(i, tx):
                for r in par.runs:
                    r.bold = True
                    r.font.color.rgb = _cor(SEGURANCA)
        return linha

    def tabela(self, cabecalhos, linhas, larguras_cm, legenda=None, primeira_negrito=True):
        if legenda:
            self.legenda(legenda)
        t = self._tabela(cabecalhos, larguras_cm)
        for l in linhas:
            self._linha(t, l, larguras_cm, primeira_negrito=primeira_negrito)
        self.doc.add_paragraph()
        return t

    def legenda(self, texto):
        par = self.doc.add_paragraph(style="Nota PCMD")
        par.paragraph_format.keep_with_next = True
        self._texto(par, texto, "parágrafo")

    def regiao(self, rid, larguras_cm, legenda, destaque=None, cabecalhos=None):
        """Tabela repetível. Template: uma linha modelo com marcadores. Exemplo: uma linha por registro."""
        spec = E.REGIOES_DOC[rid]
        campos = [c for c, _ in spec["campos"]]
        cab = cabecalhos or [rot for _, rot in spec["campos"]]
        self.legenda(legenda + (f" · região repetível «{rid}»: repita a linha modelo para cada registro." if self.valores is None else ""))
        t = self._tabela(cab, larguras_cm)
        indice = len(self.doc.tables) - 1
        if self.valores is None:
            self._linha(t, ["{{" + c + "}}" for c in campos], larguras_cm, local=f"célula de tabela (região {rid})")
        else:
            regs = self.regioes.get(rid, [])
            if not regs:
                self._linha(t, ["Nenhum registro"] + [""] * (len(campos) - 1), larguras_cm)
            for reg in regs:
                self._linha(t, [str(reg.get(c, "")) for c in campos], larguras_cm, destaque=destaque)
        self.regioes_usadas.append({"regiao": rid, "tabela_indice": indice, "linha_modelo": 1, "campos": campos})
        self.doc.add_paragraph()
        return t

    def registros(self, origem):
        """Seção 6 do modelo: dados de controle do documento e registros de origem."""
        self.h1("Registros e anexos")
        pares = [
            ("Situação", "{{documento.situacao}}", "Próxima revisão", "{{documento.proxima_revisao}}"),
            ("Origem dos dados", "{{origem.dados}}", "Data de emissão", "{{emissao.data}}"),
        ]
        larg = [3.0, 5.5, 3.0, 5.5]
        t = self.doc.add_table(rows=0, cols=4)
        tabela_fixa(t, larg)
        bordas_tabela(t)
        for a, b, c, d in pares:
            linha = t.add_row()
            nao_quebrar(linha)
            for i, tx in enumerate((a, b, c, d)):
                cel = linha.cells[i]
                cel.width = Cm(larg[i])
                par = cel.paragraphs[0]
                par.style = self.doc.styles["Tabela PCMD"]
                if i % 2 == 0:
                    sombrear(cel, CAB_FUNDO)
                    rr = par.add_run(tx)
                    rr.font.color.rgb, rr.bold = _cor(AZUL), True
                else:
                    self._texto(par, tx, "célula de tabela (identificação)")
        self.doc.add_paragraph()
        self.tabela(["Registro", "Onde fica"], origem, [6.0, 11.0])
        p = self.doc.add_paragraph()
        r = p.add_run(f"Modelo PCM Direto {VERSAO_MODELO} · {AUTOR} · {CONTATO} · Licença {LICENCA}")
        r.font.size, r.font.color.rgb = Pt(8), _cor("8C96A0")

    def secao_paisagem(self):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._pagina(sec, paisagem=True)
        # Cabeçalho e rodapé ficam vinculados à seção anterior: o aviso aparece também aqui.
        sec.header.is_linked_to_previous = True
        sec.footer.is_linked_to_previous = True
        return sec

    def checklist(self, itens):
        larg = [1.0, 11.5, 4.5]
        t = self._tabela(["Nº", "Item de conferência", "Conferido por / data"], larg)
        for i, item in enumerate(itens, start=1):
            self._linha(t, [str(i), item, ""], larg)
        self.doc.add_paragraph()

    def salvar(self, caminho):
        self.doc.core_properties.author = "PCM Direto (gerador de modelos)"
        self.doc.core_properties.title = f"{self.codigo} {self.titulo}"
        self.doc.core_properties.subject = AVISO
        self.doc.core_properties.created = dt.datetime(2026, 10, 23, 12, 0, 0)
        self.doc.core_properties.modified = dt.datetime(2026, 10, 23, 12, 0, 0)
        self.doc.core_properties.last_modified_by = "PCM Direto"
        self.doc.save(caminho)


# ------------------------------------------------------------ conteúdo comum
TERMOS = [
    ("Ativo local", "Equipamento efetivo da organização, com ID técnico, setor, TAG, plano, criticidade e histórico."),
    ("Modelo de máquina compartilhado", "Referência reutilizável do catálogo público do PCM Direto, sujeita a revisão. Usá-lo cria um cadastro local, "
                                        "sem sincronização automática."),
    ("Template documental", "Arquivo Excel/Word preenchido com dados locais, como este manual. Não consulta o catálogo nem envia dados."),
    ("ID técnico", "Identificador estável de ativo, tarefa, pessoa, execução ou falha. Não é substituído por TAG, nome ou número de linha."),
    ("TAG", "Rótulo legível do ativo no formato SETOR-SIGLA-NNN (ex.: UTL-CMP-001)."),
    ("Tarefa de segurança", "Tarefa que não pode ser retirada do plano nem ter a frequência reduzida em relação à referência."),
    ("Pendente / A definir", "Pendente: falta informação para concluir. A definir: valor ainda não definido pelo responsável; nunca é presumido."),
]


def referencias(m):
    m.p("Referencial adotado: {{organizacao.referencial}} (o inicial do PCM Direto é GERAL (sem norma)).")
    m.destaque("Cada referência começa Não verificado. Verificado exige fonte oficial HTTPS, versão e data; registra a conferência do responsável "
               "e não certifica conformidade. Não cite número de norma sem essas informações.")
    m.regiao("requisitos", [5.2, 2.0, 2.3, 4.7, 2.8], "Referências registradas (PCMD-01, aba Requisitos)")


def responsabilidades(m, papeis):
    m.h1("Responsabilidades")
    m.tabela(["Papel", "Responsabilidade", "Pessoa"], papeis, [4.2, 8.8, 4.0])
    m.regiao("pessoas", [4.0, 4.4, 3.4, 5.2], "Pessoas ativas (PCMD-01, aba Pessoas)")
    m.nota("Inclua apenas os dados pessoais necessários. Pessoas inativadas continuam associadas aos registros anteriores.")


def historico(m):
    m.h1("Histórico de revisões")
    m.regiao("revisoes", [2.0, 2.8, 8.2, 4.0], "Revisões deste documento (PCMD-01, aba Revisões)")


# ------------------------------------------------------------ W1
def manual_w1(valores=None, regioes=None):
    m = Manual("PCMD-W1", "Manual de Gestão de Ativos", valores, regioes)

    m.h1("Objetivo")
    m.p("Este manual descreve como {{organizacao.nome}} cadastra, identifica, classifica e mantém o histórico dos seus ativos de manutenção, "
        "com os registros do PCM Direto e os modelos PCMD.")

    m.h1("Escopo e aplicação")
    m.p("Inclui cadastro e hierarquia, IDs e TAGs, setores, criticidade, responsáveis, histórico, requisitos e controle documental. "
        "Não inclui procedimentos de execução, indicadores sem regra aprovada nem declaração de conformidade legal.")
    referencias(m)

    m.h1("Definições e siglas")
    m.tabela(["Termo", "Definição"], TERMOS, [4.5, 12.5])

    responsabilidades(m, [
        ("Elaborador", "Prepara e atualiza este manual e os cadastros de referência.", "{{documento.elaborador}}"),
        ("Aprovador", "Aprova o manual revisado. Preenchido só após aprovação efetiva.", "{{documento.aprovador}}"),
        ("Responsável pela criticidade", "Define o significado dos níveis 1, 2 e 3 e confere as notas.", "A definir"),
        ("Responsável por tarefa", "Pessoa ligada a cada tarefa pelo ID; renomear a pessoa não quebra o vínculo.", "Ver PCMD-03"),
    ])

    m.h1("Descrição do procedimento")
    m.h2("Cadastro, IDs e TAGs")
    m.p("Hierarquia: organização → setor → ativo → tarefas do plano. Cada ativo pertence a um setor e a um tipo de equipamento da biblioteca técnica.")
    m.tabela(["Item", "Regra"], [
        ("ID técnico", "Criado no cadastro e estável: não muda quando a descrição ou a TAG mudam."),
        ("TAG", "SETOR-SIGLA-NNN: setor (2 a 4 letras ou números), sigla de 3 letras do tipo e número de 3 dígitos."),
        ("Vínculos", "Sempre pelo ID. TAG, nome, número de linha ou título não substituem IDs."),
        ("Histórico", "Execuções e falhas guardam a TAG e a descrição do momento do registro."),
    ], [3.5, 13.5])
    m.regiao("setores", [3.0, 14.0], "Setores cadastrados (PCMD-02, aba Setores)")

    m.h2("Criticidade")
    m.p("Cinco critérios, nesta ordem: Segurança (N1), Qualidade (N2), Produção (N3), Frequência de falha (N4) e Custo/tempo de reparo (N5). "
        "Cada nota é um inteiro de 1 a 3; nota ausente é ausência, não zero.")
    m.tabela(["Situação", "Classe"], [
        ("Alguma nota ausente ou inválida", "Pendente"),
        ("Segurança = 3", "A"),
        ("Soma ≥ limite A ({{criticidade.limite_a}})", "A"),
        ("Soma ≥ limite B ({{criticidade.limite_b}})", "B"),
        ("Demais casos com cinco notas válidas", "C"),
        ("Limites fora de 5 ≤ B < A ≤ 15 ou não inteiros", "Pendente"),
    ], [12.0, 5.0], "Regra de classificação, avaliada de cima para baixo (soma de 5 a 15; limites iniciais A = 12 e B = 9)")
    m.tabela(["Critério", "Nível 1", "Nível 2", "Nível 3"], [
        (c, "A definir pelo mantenedor", "A definir pelo mantenedor", "A definir pelo mantenedor")
        for c in ("N1 Segurança", "N2 Qualidade", "N3 Produção", "N4 Frequência de falha", "N5 Custo/tempo de reparo")
    ], [4.4, 4.2, 4.2, 4.2], "Significado dos níveis (definido pelo responsável; o PCM Direto não define)")
    m.destaque("A classe não altera frequências, tarefas nem proteções de segurança. Mudanças no plano são decisão registrada do profissional responsável.")

    m.h2("Histórico dos ativos")
    m.tabela(["Situação", "O que acontece"], [
        ("Várias realizações da mesma tarefa", "Cada execução é um registro com ID próprio."),
        ("Execução cancelada", "Continua no histórico e sai da contagem do realizado."),
        ("Marcação antiga migrada", "Aparece como legado sem vínculo e não multiplica realizações."),
        ("Ativo removido", "Execuções e falhas continuam, com TAG e descrição preservadas."),
        ("Horas de parada vazias", "Significam não informado; não viram zero."),
    ], [5.5, 11.5])

    m.h2("Controle documental")
    m.p("Os documentos PCMD ficam na lista mestra (PCMD-01) com revisão, datas, elaborador, aprovador e situação; cada revisão entra no histórico.")
    m.tabela(["Código", "Documento"], [(c, t) for c, t, _ in E.DOCUMENTOS_PCMD], [3.0, 14.0])
    m.destaque("Ao usar um modelo do catálogo colaborativo, cria-se um ativo local; não há sincronização automática. Pessoas, organização, falhas e "
               "execuções são locais e não devem ser publicadas. Este manual e as planilhas não consultam a internet nem enviam dados.")

    m.h2("Indicadores")
    m.tabela(["Indicador", "Regra explícita", "Meta"], [
        ("Criticidade concluída", "Ativos com classe A, B ou C ÷ ativos cadastrados (PCMD-02).", "A definir"),
        ("Ativos Pendente", "Contagem de ativos com classe Pendente (PCMD-02).", "A definir"),
        ("Requisitos verificados", "Requisitos Verificado ÷ requisitos cadastrados (PCMD-01).", "A definir"),
        ("Disponibilidade, MTBF e MTTR", "Não calculados: regras técnicas a definir pelo mantenedor.", "{{metas.disponibilidade}}"),
    ], [4.2, 9.6, 3.2])

    m.h2("Checklist de conferência")
    m.checklist([
        "Todo ativo tem ID, TAG no formato SETOR-SIGLA-NNN, setor e tipo; não há IDs repetidos.",
        "O significado dos níveis 1, 2 e 3 foi definido; limites A e B respeitam 5 ≤ B < A ≤ 15.",
        "Ativos Pendente estão identificados e têm responsável pela avaliação.",
        "Referências Verificado têm fonte HTTPS, versão e data; nenhuma norma citada sem fonte.",
        "Histórico de execuções e falhas preservado após mudanças no cadastro.",
        "Aprovador preenchido só após aprovação efetiva; lista mestra e revisões atualizadas.",
    ])

    m.registros([
        ("Ativos e criticidade", "PCMD-02 (abas Ativos, Critérios e Setores); resumo no Anexo A"),
        ("Pessoas, requisitos e revisões", "PCMD-01 (abas Pessoas, Requisitos, Lista Mestra e Revisões)"),
        ("Execuções e falhas", "PCMD-03-04 (aba Execuções) e PCMD-05 (Ficha de Falha)"),
    ])
    historico(m)

    m.secao_paisagem()
    m.h1("Anexo A · Ativos e criticidade", numerar=False)
    m.p("Origem: PCMD-02, aba Ativos. Nota vazia = não avaliada; a pontuação só aparece com as cinco notas válidas.")
    m.regiao("ativos", [2.6, 6.4, 1.4, 3.4, 1.1, 1.1, 1.1, 1.1, 1.1, 2.4, 2.0], "Ativos e criticidade",
             destaque=lambda i, tx: i == 10 and tx == "Pendente")
    return m


# ------------------------------------------------------------ W2
def manual_w2(valores=None, regioes=None):
    m = Manual("PCMD-W2", "Manual de Estratégia de Manutenção", valores, regioes)

    m.h1("Objetivo")
    m.p("Este manual descreve como {{organizacao.nome}} escolhe e revisa estratégias de manutenção, planeja, registra execuções e trata falhas, "
        "com os registros do PCM Direto e os modelos PCMD.")

    m.h1("Escopo e aplicação")
    m.p("Aplica-se aos ativos cadastrados no PCMD-02 e às tarefas do plano mestre (PCMD-03). Cadastro, criticidade e histórico seguem o "
        "Manual de Gestão de Ativos (PCMD-W1).")
    m.destaque("Estratégias, prazos por classe e metas são definidos pelo profissional responsável. Enquanto não forem, aparecem como A definir: "
               "este manual não cria regras técnicas.")
    referencias(m)

    m.h1("Definições e siglas")
    m.tabela(["Termo", "Definição"], [
        ("Ocorrência prevista", "Trio ativo + tarefa + data prevista gerado a partir do plano vigente."),
        ("Execução vinculada", "Realização que informa a data prevista da ocorrência que atendeu."),
        ("Estratégia de manutenção", "Forma de atuação escolhida para um ativo ou tarefa pelo profissional responsável."),
        ("Tarefa de segurança", "Não pode ser retirada do plano nem ter a frequência reduzida em relação à referência."),
    ], [4.5, 12.5])
    m.nota("Demais termos (ativo local, ID técnico, TAG, Pendente, A definir): ver PCMD-W1.")

    responsabilidades(m, [
        ("Elaborador", "Prepara e atualiza este manual.", "{{documento.elaborador}}"),
        ("Aprovador", "Aprova o manual revisado. Preenchido só após aprovação efetiva.", "{{documento.aprovador}}"),
        ("Responsável pela estratégia", "Avalia consequências, escolhe e revisa estratégias e define prazos por classe.", "{{estrategia.responsavel}}"),
        ("Responsável por tarefa", "Pessoa ligada a cada tarefa do plano mestre (PCMD-03).", "Ver PCMD-03"),
    ])

    m.h1("Descrição do procedimento")
    m.h2("Estratégia por classe")
    m.p("A avaliação parte da criticidade do PCMD-02 (classe A, B, C ou Pendente). Ativos Pendente são avaliados antes de qualquer decisão baseada "
        "na classe, e cada decisão tem justificativa registrada.")
    m.tabela(["Classe", "Estratégia predominante", "Prazo de atendimento", "Revisão do plano"], [
        (c.upper(), f"{{{{estrategia.classe_{c}.estrategia}}}}", f"{{{{estrategia.classe_{c}.prazo}}}}", f"{{{{estrategia.classe_{c}.revisao}}}}")
        for c in "abc"
    ] + [("Pendente", "Avaliar a criticidade antes de definir", "Não se aplica", "Não se aplica")],
        [2.2, 5.6, 4.6, 4.6], "Definição do profissional responsável ({{estrategia.responsavel}})")
    m.p("Periodicidade de revisão da estratégia: {{estrategia.periodicidade_revisao}}. A classe não altera automaticamente nenhuma frequência.")

    m.h2("Origem das tarefas")
    m.tabela(["Regra", "Detalhe"], [
        ("Origem e versão", "Toda tarefa vem da biblioteca técnica com origem e versão registradas."),
        ("Tarefa genérica", "Origem exata: base genérica a validar com manual do fabricante."),
        ("Confirmação", "Ao confirmar em manual do fabricante ou outra fonte, atualizar origem e versão."),
    ], [4.0, 13.0])
    m.destaque("Tarefa de segurança não pode ser excluída, desmarcada nem ter intervalo ampliado; só é permitido manter ou aumentar a frequência. "
               "Validações de planilha ajudam, mas podem ser contornadas e não substituem as travas do app.")

    m.h2("Do plano ao registro")
    m.tabela(["Etapa", "O que fazer", "Onde"], [
        ("1. Planejar", "Manter atividade, frequência vigente, especialidade, duração e responsável de cada tarefa.", "PCMD-03 Plano Mestre"),
        ("2. Prever", "Gerar as ocorrências do ciclo de 52 semanas a partir do início (sem calendário ISO; diária = dias úteis, sem feriados).",
         "PCMD-04 Ocorrências"),
        ("3. Executar", "Registrar cada realização com data, responsável e a data prevista da ocorrência atendida.", "PCMD-04 Execuções"),
        ("4. Acompanhar", "Ver o % realizado por mês da data prevista e comparar com a meta {{metas.aderencia}}.", "PCMD-04 Cronograma"),
    ], [2.6, 10.4, 4.0])
    m.p("Como o realizado é contado:")
    m.lista([
        "Pelo vínculo explícito à ocorrência prevista, nunca pela coincidência da data real; o mês é o da data prevista.",
        "Duas execuções da mesma ocorrência contam uma vez. Canceladas e sem vínculo (inclusive legado) aparecem separadas e não contam.",
        "Horas previstas são estimadas pela duração da tarefa; não são horas trabalhadas.",
    ])

    m.h2("Tratamento de falhas")
    m.p("Cada falha vai para a ficha PCMD-05 com ativo, data, sintoma, causa, ação, horas de parada (opcional), peça e responsável. "
        "A análise de causa segue o critério do profissional responsável; este manual não prescreve procedimentos técnicos.")
    m.lista([
        "Horas de parada vazias significam não informado e não são tratadas como zero.",
        "Falhas recorrentes vão para a revisão da estratégia, na periodicidade definida.",
    ])

    m.h2("Indicadores")
    m.tabela(["Indicador", "Regra explícita", "Meta"], [
        ("% realizado (por ocorrência)", "Concluídas distintas ÷ previstas distintas, no mês da data prevista; mês sem previsão: Sem previsão.", "{{metas.aderencia}}"),
        ("Falhas registradas", "Contagem de registros da PCMD-05 no período.", "A definir"),
        ("Horas de parada informadas", "Soma das horas informadas na PCMD-05; vazios não entram como zero. Não é disponibilidade.", "A definir"),
        ("Disponibilidade, MTBF e MTTR", "Não calculados: regras técnicas a definir pelo mantenedor.", "{{metas.disponibilidade}}"),
    ], [4.2, 9.6, 3.2])

    m.h2("Revisão da estratégia")
    m.p("A estratégia é revisada na periodicidade {{estrategia.periodicidade_revisao}} e sempre que mudarem o ativo, a criticidade, o histórico "
        "de falhas ou as referências. Cada revisão entra no histórico abaixo e na lista mestra (PCMD-01).")

    m.h2("Checklist de conferência")
    m.checklist([
        "Estratégia por classe definida pelo responsável ou mantida como A definir; Pendente avaliado antes.",
        "Tarefas de segurança com frequência igual ou maior que a referência e não retiradas.",
        "Toda tarefa tem origem e versão; genéricas trazem a origem exata da base genérica.",
        "Execuções informam a data prevista; canceladas e sem vínculo fora do percentual.",
        "Falhas com ativo, data e sintoma; horas de parada vazias não tratadas como zero.",
        "Metas sem regra aprovada como A definir; aprovador só após aprovação efetiva.",
    ])

    m.registros([
        ("Plano mestre", "PCMD-03-04 (aba Plano Mestre); resumo no Anexo A"),
        ("Ocorrências, execuções e cronograma", "PCMD-03-04 (abas Ocorrências, Execuções e Cronograma)"),
        ("Falhas", "PCMD-05 (aba Ficha de Falha)"),
    ])
    historico(m)

    m.secao_paisagem()
    m.h1("Anexo A · Plano mestre (resumo)", numerar=False)
    m.p("Origem: PCMD-03, aba Plano Mestre. Segurança = Sim: a tarefa não pode ser retirada nem ter a frequência reduzida.")
    m.regiao("tarefas", [2.4, 1.7, 6.6, 2.2, 2.7, 2.0, 2.3, 3.4, 2.4], "Tarefas do plano",
             destaque=lambda i, tx: i == 6 and tx == "Sim")
    return m


# ------------------------------------------------------------ dados do exemplo
def valores_exemplo(d):
    g = d["gestao"]
    pessoas = {p["id"]: p["nome"] for p in g["pessoas"]}
    org = g["organizacao"]
    v = {
        "organizacao.nome": org["nome"], "organizacao.referencial": org["referencial"],
        "documento.revisao": "00", "documento.data": "23/10/2026",
        "documento.elaborador": pessoas.get(org["elaborador"], ""),
        "documento.aprovador": pessoas.get(org["aprovador"], "") or "A definir (aprovação não registrada)",
        "documento.proxima_revisao": "", "documento.situacao": "Em elaboração",
        "emissao.data": "23/10/2026", "origem.dados": "Demonstração fictícia (não é exportação do app)",
        "criticidade.limite_a": g["criticidade"]["limites"]["a"], "criticidade.limite_b": g["criticidade"]["limites"]["b"],
        "metas.disponibilidade": "" if g["metas"]["disponibilidade"] is None else f"{g['metas']['disponibilidade']}%",
        "metas.aderencia": "" if g["metas"]["aderencia"] is None else f"{g['metas']['aderencia']}%",
    }
    with open(RAIZ_REPO / "src" / "dados" / "biblioteca.json", encoding="utf-8") as f:
        nomes_tipo = {t["id"]: t["nome"] for t in json.load(f)["tipos"]}
    status = {"verificado": "Verificado", "nao_verificado": "Não verificado"}
    reg = {
        "requisitos": [{"requisito.referencia": q["referencia"], "requisito.versao": q["versao"] or "A definir",
                        "requisito.data": q["data"] or "A definir", "requisito.fonte": q["fonte"] or "A definir",
                        "requisito.status": status[q["status"]]} for q in g["requisitos"]],
        "pessoas": [{"pessoa.nome": p["nome"], "pessoa.funcao": p["funcao"], "pessoa.especialidade": p["especialidade"],
                     "pessoa.responsabilidade": "A definir"} for p in g["pessoas"] if p["ativo"]],
        "setores": [{"setor.sigla": s, "setor.descricao": "A definir"} for s in org["setores"]],
        "revisoes": [{"revisao.numero": "00", "revisao.data": "23/10/2026", "revisao.descricao": "Emissão inicial (exemplo fictício)",
                      "revisao.responsavel": pessoas.get(org["elaborador"], "")}],
    }
    ativos = []
    for e in d["equipamentos"]:
        n = g["criticidade"]["notas"].get(e["id"], [None] * 5)
        c = d["criticidade"].get(e["id"], {"pontuacao": None, "classe": "Pendente"})
        ativos.append({"ativo.tag": e["tag"], "ativo.descricao": e["descricao"], "ativo.setor": e["setor"], "ativo.tipo_nome": nomes_tipo.get(e["tipo"], e["tipo"]),
                       "criticidade.nota_seguranca": n[0] or "", "criticidade.nota_qualidade": n[1] or "", "criticidade.nota_producao": n[2] or "",
                       "criticidade.nota_frequencia_falha": n[3] or "", "criticidade.nota_custo_reparo": n[4] or "",
                       "criticidade.pontuacao": c["pontuacao"] if c["pontuacao"] is not None else "", "criticidade.classe": c["classe"]})
    reg["ativos"] = ativos
    tarefas = []
    for l in sorted(d["plano"], key=lambda x: (x["tag"], x["tarefaId"])):
        resp = g["responsaveis"].get(f"{l['equipamentoId']}|{l['tarefaId']}", "")
        tarefas.append({"tarefa.tag": l["tag"], "tarefa.id": l["tarefaId"], "tarefa.atividade": l["descricao"], "tarefa.frequencia_vigente": l["periodicidade"],
                        "tarefa.especialidade": l["especialidade"], "tarefa.duracao_horas": f"{l['duracaoHoras']:.2f}".replace(".", ","),
                        "tarefa.seguranca": "Sim" if l["seguranca"] else "Não", "tarefa.origem_versao": f"{l['origem']} · {l['versao']}",
                        "tarefa.responsavel_nome": pessoas.get(resp, "")})
    reg["tarefas"] = tarefas
    return v, reg
