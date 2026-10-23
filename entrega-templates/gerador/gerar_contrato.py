"""Gera CONTRATO-CAMPOS.json e DICIONARIO-DE-DADOS.md a partir da especificação e dos arquivos gerados."""
import json
from collections import OrderedDict

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as letra

import especificacao as E
from comum import AVISO, DATA_MODELO, RAMO_GESTAO, VERSAO_MODELO

TIPOS = {
    "texto": "texto", "lista": "texto de lista fechada", "sim_nao": "texto 'Sim' ou 'Não'", "data": "data do Excel (número de série)",
    "numero_horas": "número decimal de horas", "inteiro": "número inteiro", "percentual": "fração 0–1 com formato %", "booleano": "VERDADEIRO/FALSO",
}


def campo_coluna(col, arquivo, aba, indice, l1, l2):
    L = letra(indice)
    return OrderedDict([
        ("id", col.id), ("rotulo", col.cabecalho), ("tipo", TIPOS.get(col.tipo, col.tipo)), ("obrigatorio", col.obrigatorio),
        ("formato_unidade", col.formato if col.formato != "texto" else ("texto (célula com formato Texto)" if col.papel != "calculo" else "texto")),
        ("papel", col.papel), ("entrada_ou_calculo", "cálculo" if col.papel == "calculo" else "entrada"),
        ("origem_proposta", col.origem), ("situacao_origem", col.origem_status),
        ("destino", {"arquivo": arquivo, "aba": aba, "coluna": L, "cabecalho_linha": E.LINHA_CAB, "intervalo_dados": f"{L}{l1}:{L}{l2}",
                     "intervalo_nomeado": col.nome}),
        ("regra_vazio", col.vazio), ("validacao", col.validacao if col.validacao != "—" else (col.dv["type"] if col.dv else "—")),
        ("validacao_excel", None if not col.dv else {k: v for k, v in col.dv.items() if k in ("type", "operator", "formula1", "formula2")}),
        ("formula_modelo", col.formula), ("privacidade", col.privacidade), ("observacoes", col.obs),
    ])


LISTAS = {"tipo_documento": ("lst_tipo_documento", E.TIPOS_DOCUMENTO), "situacao_documento": ("lst_situacao_documento", E.SITUACOES_DOCUMENTO),
          "sim_nao": ("lst_sim_nao", ["Sim", "Não"]), "status_requisito": ("lst_status_requisito", ["Não verificado", "Verificado"]),
          "frequencias": ("lst_frequencias", [f[0] for f in E.FREQUENCIAS]), "origem_execucao": ("lst_origem_execucao", ["app", "legado"]),
          "especialidades": ("lst_especialidades", "biblioteca.tipos[].tarefas[].especialidade (sem repetição)"),
          "tipos": ("lst_tipo_id", "biblioteca.tipos[].id (com lst_tipo_nome e lst_tipo_sigla)"),
          "tipos_tarefa": (None, E.TIPOS_TAREFA_REFERENCIA)}


def gerar(destino_json, destino_md, xlsx, docs):
    """xlsx: lista de (Arquivo, caminho do modelo em branco); docs: lista de (nome do arquivo, Manual)."""
    campos_todos = OrderedDict()
    arquivos_xlsx = []

    def registrar(c):
        campos_todos.setdefault(c["id"], []).append(c["destino"])

    for arq, caminho in xlsx:
        nome_arq = f"planilhas/{arq.nome}.xlsx"
        wb = load_workbook(caminho)
        nomes = [{"nome": n, "referencia": d.attr_text} for n, d in sorted(wb.defined_names.items())]
        abas = []
        for aba in arq.tabelas:
            l1, l2 = E.PRIMEIRA, E.PRIMEIRA + aba.capacidade - 1
            cols = [campo_coluna(c, nome_arq, aba.nome, i + 1, l1, l2) for i, c in enumerate(aba.colunas)]
            for c in cols:
                registrar(c)
            abas.append(OrderedDict([("nome", aba.nome), ("codigo", aba.codigo), ("tipo_regiao", "tabela de linhas repetíveis"),
                                     ("linha_papel", E.LINHA_TIPO), ("linha_cabecalho", E.LINHA_CAB), ("primeira_linha", l1),
                                     ("ultima_linha_com_formulas", l2), ("capacidade_linhas", aba.capacidade), ("protegida_sem_senha", aba.protegida),
                                     ("colunas", cols)]))
        celulas = []
        for (cid, rot, aba, cel, papel, tipo, origem, sit, vazio, val, priv, obs) in E.CAMPOS_CELULAS:
            if aba not in wb.sheetnames:
                continue
            d = OrderedDict([("id", cid), ("rotulo", rot), ("tipo", TIPOS.get(tipo, tipo)), ("papel", papel),
                             ("entrada_ou_calculo", "cálculo" if papel == "calculo" else "entrada"), ("origem_proposta", origem),
                             ("situacao_origem", sit), ("destino", {"arquivo": nome_arq, "aba": aba, "celula": cel}), ("regra_vazio", vazio),
                             ("validacao", val), ("privacidade", priv), ("observacoes", obs)])
            celulas.append(d)
            registrar(d)
        for cid, rot, nm, tipo in (("emissao.origem_dados", "Origem dos dados", "pcmd_origem", "texto"), ("emissao.data", "Data de emissão", "pcmd_emissao", "data")):
            cel = wb.defined_names[nm].attr_text.split("!")[1].replace("$", "")
            d = OrderedDict([("id", cid), ("rotulo", rot), ("tipo", TIPOS[tipo]), ("papel", "entrada"), ("entrada_ou_calculo", "entrada"),
                             ("origem_proposta", "Texto curto do adaptador ou de quem preenche" if tipo == "texto" else "Data em que o arquivo é gerado/preenchido"),
                             ("situacao_origem", E.FUTURO), ("destino", {"arquivo": nome_arq, "aba": "Início", "celula": cel, "intervalo_nomeado": nm}),
                             ("regra_vazio", "Vazio = não informado."), ("validacao", "Texto." if tipo == "texto" else "Data válida."),
                             ("privacidade", E.PRIV_NENHUM), ("observacoes", "Existe em cada arquivo XLSX.")])
            celulas.append(d)
            registrar(d)
        listas = OrderedDict((LISTAS[k][0] or k, LISTAS[k][1]) for k in arq.listas)
        item = OrderedDict([("arquivo", nome_arq), ("exemplo", f"planilhas/{arq.nome}-Exemplo.xlsx"), ("codigo", arq.codigo), ("titulo", arq.titulo),
                            ("abas", wb.sheetnames), ("abas_tabela", abas), ("celulas_unicas", celulas), ("intervalos_nomeados", nomes),
                            ("listas", listas)])
        if "Cronograma" in wb.sheetnames:
            cr = E.CRONO
            item["cronograma_resumo"] = OrderedDict([
                ("descricao", "Resumo mensal por ocorrência (equivale a resumoMensal do ramo de gestão). Somente cálculo; o adaptador não escreve aqui."),
                ("cartoes", f"linha {cr['cartao']}: previstas (A), concluídas (C), % realizado (D), meta (E), situação frente à meta (G)"),
                ("meses", f"linhas {cr['jan']} a {cr['jan'] + 11}; total na linha {cr['total']}"),
                ("colunas", {"C": "previstas distintas = SUMIFS(oc_conta, oc_ano, crono_ano, oc_mes, mês)",
                             "D": "concluídas distintas = SUMIFS(oc_concluida, oc_ano, crono_ano, oc_mes, mês)",
                             "E": "% = D/C; 'Sem previsão' quando C = 0 (app: percentual null; app usa escala 0–100, aqui fração 0–1)",
                             "F": "horas estimadas = SUMIFS(oc_horas, oc_conta, 1, oc_ano, crono_ano, oc_mes, mês)"}),
                ("grafico", f"colunas previstas × concluídas por mês, ancorado em A{cr['grafico']}"),
                ("fora_do_ano", f"linhas {cr['fora_cab'] + 1} (Antes do ano) e {cr['fora_cab'] + 2} (Depois do ano)"),
                ("sem_vinculo", f"linhas {cr['semv_cab'] + 1} a {cr['semv_cab'] + 3}: sem vínculo, previsão não encontrada, canceladas"),
                ("conferencias", f"linhas {cr['conf_cab'] + 1} a {cr['conf_cab'] + 4}"),
            ])
        arquivos_xlsx.append(item)
    arquivos_docx = []
    marcadores_doc = {m[0]: m for m in E.MARCADORES_DOC}
    for nome_arq, manual in docs:
        locais = OrderedDict()
        for mk, local in manual.marcadores_usados:
            locais.setdefault(mk, set()).add(local)
        simples = []
        for mk, ls in locais.items():
            if mk in marcadores_doc:
                _, desc, tipo, origem, sit, vazio = marcadores_doc[mk]
                simples.append(OrderedDict([("marcador", "{{" + mk + "}}"), ("id", mk), ("descricao", desc), ("tipo", TIPOS.get(tipo, tipo)),
                                            ("locais", sorted(ls)), ("origem_proposta", origem), ("situacao_origem", sit), ("regra_vazio", vazio)]))
        regioes = []
        for r in manual.regioes_usadas:
            spec = E.REGIOES_DOC[r["regiao"]]
            regioes.append(OrderedDict([("regiao", r["regiao"]), ("descricao", spec["descricao"]), ("tabela_indice_no_corpo", r["tabela_indice"]),
                                        ("linha_cabecalho", 0), ("linha_modelo", r["linha_modelo"]),
                                        ("marcadores_da_linha", ["{{" + c + "}}" for c in r["campos"]]),
                                        ("origem_proposta", spec["origem"]), ("situacao_origem", spec["situacao"]),
                                        ("regra", "Repetir a linha modelo uma vez por registro; sem registros, escrever 'Nenhum registro' na primeira célula.")]))
        arquivos_docx.append(OrderedDict([("arquivo", nome_arq), ("codigo", manual.codigo), ("titulo", manual.titulo),
                                          ("marcadores", simples), ("regioes_repetiveis", regioes)]))
        for s in simples:
            campos_todos.setdefault("docx:" + s["id"], []).append({"arquivo": nome_arq, "locais": s["locais"]})
    contrato = OrderedDict([
        ("$comentario", "Mapa PROPOSTO para um adaptador futuro. Não é formato de backup nem API aceita pelo PCM Direto. Não contém dados reais."),
        ("contrato", "PCM Direto · Contrato de campos dos modelos PCMD"),
        ("versao", 1), ("versao_modelo", VERSAO_MODELO), ("data", DATA_MODELO),
        ("situacao", "proposto, ainda não validado por adaptador ou pelo app"),
        ("aviso", AVISO),
        ("fontes_consultadas", [
            "AGENTS.md, BIBLIOTECA.md, src/marca.js (main, commit 6f09c92)",
            "src/logica/plano.js, src/logica/cronograma.js, src/logica/tag.js, src/logica/pacote.js, src/estado.js, src/dados/biblioteca.json (main)",
            f"src/logica/gestao.js, src/estado.js e src/telas/cadastros.js do ramo {RAMO_GESTAO}",
            "docs/INTERFACE-E-MODELOS.md (pasta de contexto do projeto; mesmo texto do ramo acima)",
        ]),
        ("limites", [
            "dadosAtuais() (montarDados) não inclui estado.gestao; o adaptador precisará ler estado.gestao separadamente.",
            "O exportador atual (src/saida/xlsx.js e docx.js) não escreve estes modelos.",
            "estado.gestao existe apenas no ramo indicado; se o ramo mudar antes da integração, revalidar este contrato.",
            "Este contrato não autoriza upload de backup nem envio de dados; o preenchimento deve ocorrer localmente.",
        ]),
        ("convencoes", OrderedDict([
            ("chave_tarefa", "equipamentoId|tarefaId"), ("chave_ocorrencia", "equipamentoId|tarefaId|AAAA-MM-DD (data prevista)"),
            ("ids", "IDs técnicos são texto [A-Za-z0-9_-]; nunca substituir por TAG, nome ou número de linha."),
            ("datas", "App usa texto AAAA-MM-DD; no XLSX gravar data real do Excel com formato dd/mm/aaaa. '' no app → célula vazia."),
            ("horas", "Número decimal (0.5 = meia hora). null no app → célula vazia (não zero)."),
            ("percentuais", "App guarda metas e percentuais em 0–100; no XLSX gravar fração 0–1 com formato %. Percentual do resumo é cálculo."),
            ("booleanos", "true/false no app → 'Sim'/'Não' no XLSX."),
            ("textos", "Gravar como texto (tipo string); conteúdo iniciado por =, +, - ou @ nunca vira fórmula."),
            ("formulas", "Gravadas em inglês/OOXML; o adaptador não deve sobrescrever colunas de cálculo nem nomes definidos."),
            ("capacidade", "Se os dados excederem a capacidade, estender fórmulas, validações e intervalos nomeados antes de escrever (ver COMO-USAR.md)."),
        ])),
        ("situacoes_origem", E.SITUACAO_ORIGEM),
        ("papeis", {k: v[1] for k, v in E.PAPEIS.items()}),
        ("ids_tecnicos", [
            {"id": "ativo.id", "onde": "PCMD-02 Ativos A; PCMD-03/04 Plano Mestre A, Ocorrências A, Execuções B; PCMD-05 Ficha de Falha B",
             "origem": "estado.equipamentos[].id"},
            {"id": "tarefa.id", "onde": "PCMD-03/04 Plano Mestre B, Ocorrências B, Execuções C", "origem": "biblioteca.tipos[].tarefas[].id"},
            {"id": "pessoa.id", "onde": "PCMD-01 Pessoas A e Organização B10/B11; PCMD-03/04 Plano Mestre P, Execuções K; PCMD-05 Ficha de Falha J",
             "origem": "gestao.pessoas[].id"},
            {"id": "execucao.id", "onde": "PCMD-03/04 Execuções A", "origem": "gestao.historico[].id"},
            {"id": "falha.id", "onde": "PCMD-05 Ficha de Falha A", "origem": "gestao.falhas[].id"},
            {"id": "requisito.id", "onde": "PCMD-01 Requisitos A", "origem": "gestao.requisitos[].id"},
        ]),
        ("independencia_dos_arquivos", "Cada XLSX é independente: nenhuma fórmula consulta outro arquivo. Nomes de pessoas e TAGs são gravados "
                                       "como texto ao lado dos IDs; o adaptador deve escrever os dois de forma coerente."),
        ("arquivos", arquivos_xlsx + arquivos_docx),
        ("indice_campos", campos_todos),
    ])
    with open(destino_json, "w", encoding="utf-8") as f:
        json.dump(contrato, f, ensure_ascii=False, indent=2)
        f.write("\n")
    escrever_dicionario(destino_md, arquivos_xlsx, arquivos_docx)


def _esc(t):
    return str(t if t is not None else "").replace("|", "\\|").replace("\n", " ")


def escrever_dicionario(destino, xlsx, docx):
    rot_sit = {E.MAIN: "app (main)", E.RAMO: "app (ramo de gestão, não integrado)", E.MANUAL: "manual", E.FUTURO: "futuro", E.CALCULO: "cálculo"}
    L = ["# Dicionário de dados dos modelos PCMD", "",
         f"Modelo versão {VERSAO_MODELO} · gerado a partir de `gerador/especificacao.py` (mesma fonte do `CONTRATO-CAMPOS.json`).", "",
         f"> **Aviso:** {AVISO}", "",
         "**Situação do mapeamento: proposto, ainda não validado.** A coluna Origem indica de onde o futuro adaptador poderia ler cada dado. "
         f"Campos marcados como *ramo de gestão* existem apenas no ramo `{RAMO_GESTAO}`. O app ainda não preenche estes arquivos.", "",
         "Papéis: PREENCHER (entrada ou manual), ID (identificador técnico), REFERÊNCIA (protegida no Plano Mestre), AUTOMÁTICO (fórmula; não escrever), "
         "FUTURO (manter vazio).", ""]
    for x in xlsx:
        L += [f"# {x['arquivo']}", "", f"{x['codigo']} · {x['titulo']}. Abas: {', '.join(x['abas'])}. Versão com dados fictícios: `{x['exemplo']}`.", ""]
        for a in x["abas_tabela"]:
            L += [f"## {a['nome']}", "",
                  f"Cabeçalho na linha {a['linha_cabecalho']}; dados de {a['primeira_linha']} a {a['ultima_linha_com_formulas']} "
                  f"(capacidade {a['capacidade_linhas']} linhas){'; aba protegida sem senha' if a['protegida_sem_senha'] else ''}.", "",
                  "| Col. | Campo | Cabeçalho | Papel | Tipo / formato | Obrig. | Origem proposta | Situação | Vazio | Validação |",
                  "|---|---|---|---|---|---|---|---|---|---|"]
            for c in a["colunas"]:
                L.append(f"| {c['destino']['coluna']} | `{c['id']}` | {_esc(c['rotulo'])} | {E.PAPEIS[c['papel']][0]} | {_esc(c['formato_unidade'])} | "
                         f"{c['obrigatorio']} | {_esc(c['origem_proposta'])} | {rot_sit[c['situacao_origem']]} | {_esc(c['regra_vazio'])} | {_esc(c['validacao'])} |")
            L.append("")
        L += ["## Células únicas", "", "| Aba!Célula | Campo | Rótulo | Papel | Origem proposta | Situação | Vazio | Validação |", "|---|---|---|---|---|---|---|---|"]
        for c in x["celulas_unicas"]:
            L.append(f"| {c['destino']['aba']}!{c['destino']['celula']} | `{c['id']}` | {_esc(c['rotulo'])} | {E.PAPEIS[c['papel']][0]} | "
                     f"{_esc(c['origem_proposta'])} | {rot_sit[c['situacao_origem']]} | {_esc(c['regra_vazio'])} | {_esc(c.get('validacao', '—'))} |")
        L += ["", "## Intervalos nomeados", "", "| Nome | Referência |", "|---|---|"] + \
             [f"| `{n['nome']}` | `{n['referencia']}` |" for n in x["intervalos_nomeados"]] + [""]
    for d in docx:
        L += [f"# {d['arquivo']}", "", "### Marcadores", "", "| Marcador | Descrição | Locais | Origem proposta | Situação | Vazio |", "|---|---|---|---|---|---|"]
        for mk in d["marcadores"]:
            L.append(f"| `{mk['marcador']}` | {_esc(mk['descricao'])} | {', '.join(mk['locais'])} | {_esc(mk['origem_proposta'])} | "
                     f"{rot_sit[mk['situacao_origem']]} | {_esc(mk['regra_vazio'])} |")
        L += ["", "### Regiões repetíveis", "", "| Região | Tabela nº (corpo, a partir de 0) | Linha modelo | Marcadores da linha |", "|---|---|---|---|"]
        for r in d["regioes_repetiveis"]:
            L.append(f"| `{r['regiao']}` | {r['tabela_indice_no_corpo']} | {r['linha_modelo']} | {' '.join('`' + x + '`' for x in r['marcadores_da_linha'])} |")
        L.append("")
    with open(destino, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
