"""Gera os quatro arquivos XLSX dos modelos PCMD, cada um em branco e com exemplo fictício.

Cada arquivo é independente: nenhuma fórmula depende de outro arquivo. Modelo e exemplo saem da
mesma função de montagem; o exemplo apenas recebe dados. Fórmulas são gravadas em inglês (padrão
OOXML); o Excel em português mostra SE, ÍNDICE, CORRESP...
"""
import datetime as dt
import json
import math

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as letra
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.pagebreak import Break

import especificacao as E
from comum import (AUTOR, AVISO, AVISO_FUNDO, AVISO_TEXTO, AZUL, CAB_FUNDO, CINZA_APOIO, CONTATO, DESTAQUE, ENTRADA_FUNDO, ENTRADA_TEXTO,
                   FONTE, FONTE_ESTREITA, FONTE_TITULO, GRADE, LICENCA, LOGO_MARCA, RAIZ_REPO, SECAO_BORDA, SECAO_FUNDO, SEGURANCA, TEXTO,
                   VERSAO_MODELO, ZEBRA_PADRAO, dividir_aviso)

# ------------------------------------------------------------------ padrão visual PCM Direto (modelo base de planilha)
TINTA = TEXTO
CINZA = CINZA_APOIO
ZEBRA = ZEBRA_PADRAO
CALC_FUNDO = "F6F8FC"   # células automáticas e rótulos
ERRO_FUNDO = "FBE7EA"
AZUL_SUAVE = "E3ECF4"
CAB_CALC = SECAO_FUNDO  # cabeçalho das colunas automáticas (mais claro, texto cinza)


def preench(cor):
    return PatternFill("solid", fgColor=cor)


def fonte(tam=9, cor=TINTA, negrito=False, italico=False, sublinhado=None, nome_fonte=FONTE):
    return Font(name=nome_fonte, size=tam, color=cor, bold=negrito, italic=italico, underline=sublinhado)


FINO = Side(style="thin", color=GRADE)
B_GRADE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
B_LINHA = B_GRADE
B_CAIXA = B_GRADE
B_CAB = Border(left=FINO, right=FINO, top=FINO, bottom=Side(style="medium", color=AZUL))
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
FORMATO_TEXTO = "@"
AVISO_L1, AVISO_L2 = dividir_aviso()
LINHA_FAIXA = E.LINHA_TIPO  # 5
PX_PAGINA = {True: 1036, False: 707}  # largura útil da página A4 (paisagem/retrato) em pixels, com margens de 0,45"


# ------------------------------------------------------------------ utilitários
def texto(cell, valor):
    """Grava texto como texto, mesmo que comece com =, +, - ou @ (não vira fórmula)."""
    cell.value = valor
    if isinstance(valor, str):
        cell.data_type = "s"
        if valor[:1] in "=+-@":
            cell.quotePrefix = True


def nome(wb, nome_def, ref):
    wb.defined_names[nome_def] = DefinedName(nome_def, attr_text=ref)


def ref_abs(aba, col, l1, l2=None):
    c = letra(col) if isinstance(col, int) else col
    return f"'{aba}'!${c}${l1}" + (f":${c}${l2}" if l2 else "")


def largura(ws, c):
    d = ws.column_dimensions[letra(c)]
    return 0 if d.hidden else (d.width or 9)


def largura_total(ws, c1, c2):
    return sum(largura(ws, c) for c in range(c1, c2 + 1))


def px(ws, c):
    w = largura(ws, c)
    return 0 if w == 0 else int(w * 7 + 5)


def altura_texto(txt, larg, tam=9, minimo=15):
    """Altura aproximada de um texto com quebra em células mescladas (o Excel não ajusta mescladas)."""
    por_linha = max(8, larg * 1.15 * 10 / tam)
    linhas = sum(max(1, math.ceil(len(p) / por_linha)) for p in str(txt).split("\n"))
    return max(minimo, linhas * tam * 1.45 + 5)


def mesclar(ws, linha, c1, c2, linha2=None):
    linha2 = linha2 or linha
    if c2 > c1 or linha2 > linha:
        ws.merge_cells(start_row=linha, start_column=c1, end_row=linha2, end_column=c2)


def contorno(ws, r1, c1, r2, c2, lado, baixo=None):
    """Borda externa de um intervalo, preservando as demais bordas das células."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            b = ws.cell(r, c).border
            ws.cell(r, c).border = Border(left=lado if c == c1 else b.left, right=lado if c == c2 else b.right,
                                          top=lado if r == r1 else b.top, bottom=(baixo or lado) if r == r2 else b.bottom)


def configurar_impressao(ws, paisagem, ultima_col, ultima_linha, codigo, titulos=None, quebra=None, colunas=None, linha_ini=1):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.page_setup.pageOrder = "overThenDown"  # partes esquerda e direita da mesma linha em sequência
    ws.page_margins.left = ws.page_margins.right = 0.45
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.85
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3
    if quebra:  # abas largas: duas partes com quebra fixa entre grupos e escala única que cabe nas duas
        rep = range(1, ws[colunas.split(":")[1] + "1"].column + 1) if colunas else range(0)
        p1 = sum(px(ws, c) for c in range(1, quebra))
        p2 = sum(px(ws, c) for c in rep) + sum(px(ws, c) for c in range(quebra, ultima_col + 1))
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = max(40, min(100, int(PX_PAGINA[paisagem] * 0.95 / max(p1, p2) * 100)))
        ws.col_breaks.append(Break(id=quebra - 1))
    else:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    if titulos:
        ws.print_title_rows = titulos
    if colunas:
        ws.print_title_cols = colunas
    ws.print_area = f"A{linha_ini}:{letra(ultima_col)}{ultima_linha}"
    # Rodapé: o aviso completo em duas linhas (menos de 255 caracteres com os códigos). Os dados do rodapé do
    # modelo padrão (arquivo, aba, código, data de impressão e página) ficam no cabeçalho da página.
    ws.oddFooter.left.text = f"{AVISO_L1}\n{AVISO_L2}"
    ws.oddFooter.left.size = 7
    ws.oddFooter.left.color = CINZA
    for parte, txt in ((ws.oddHeader.left, "&F  ·  &A"), (ws.oddHeader.center, f"PCM Direto  ·  {codigo}"),
                       (ws.oddHeader.right, "Impresso em &D  ·  Pág. &P de &N")):
        parte.text, parte.size, parte.color, parte.font = txt, 7.5, CINZA, f"{FONTE_ESTREITA},Regular"


# ------------------------------------------------------------------ bloco de cabeçalho, aviso, faixa e cabeçalho da tabela
F_REV = ('="Rev. "&IF(cfg_revisao="","—",cfg_revisao)&IF(ISNUMBER(cfg_data_revisao),"   ·   "&RIGHT("0"&DAY(cfg_data_revisao),2)&"/"'
         '&RIGHT("0"&MONTH(cfg_data_revisao),2)&"/"&YEAR(cfg_data_revisao),"")')
F_ELAB = '="Elaborado: "&IF(cfg_elaborado="","A definir",cfg_elaborado)&"   ·   Aprovado: "&IF(cfg_aprovado="","A definir",cfg_aprovado)'


def partes_bloco(ws, c1, c2, logo_fim=None):
    """Colunas do logo, do título e da revisão dentro de c1..c2 (larguras visíveis)."""
    if logo_fim is None:
        acc, logo_fim = 0, c1
        for c in range(c1, c2 + 1):
            acc += largura(ws, c)
            logo_fim = c
            if acc >= 22:
                break
    acc, dir_ini = 0, c2
    for c in range(c2, logo_fim, -1):
        acc += largura(ws, c)
        dir_ini = c
        if acc >= 22:
            break
    if dir_ini <= logo_fim + 1 or largura_total(ws, logo_fim + 1, dir_ini - 1) < 12:
        dir_ini = c2 + 1  # pouco espaço: sem bloco de revisão
    return logo_fim, dir_ini


def logo(ws, c1, c2):
    """Assinatura com slogan centralizada nas colunas c1..c2, linhas 1-2 (proporção preservada)."""
    if not LOGO_MARCA.exists():
        return
    from openpyxl.drawing.image import Image as Imagem
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    img = Imagem(str(LOGO_MARCA))
    prop = img.width / img.height
    disp = sum(px(ws, c) for c in range(c1, c2 + 1))
    alt_px = int((ws.row_dimensions[1].height + ws.row_dimensions[2].height) * 96 / 72)
    w = min(disp - 14, int((alt_px - 16) * prop), 250)
    h = int(w / prop)
    x, col = (disp - w) // 2, c1
    while col < c2 and x >= px(ws, col):
        x -= px(ws, col)
        col += 1
    y = (alt_px - h) // 2
    emu = 9525
    img.width, img.height = w, h
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=col - 1, colOff=x * emu, row=0, rowOff=y * emu), ext=XDRPositiveSize2D(w * emu, h * emu))
    ws.add_image(img)


def bloco(ws, codigo, titulo, c1, c2, logo_fim=None, continuacao=False):
    """Linhas 1-2: logo | código e título | revisão e elaborado/aprovado (lidos da aba Config)."""
    ws.row_dimensions[1].height = ws.row_dimensions[2].height = 33
    azul = Side(style="thin", color=AZUL)
    base = Side(style="medium", color=AZUL)
    if continuacao:
        logo_fim, dir_ini = c1 - 1, c2 + 1
    else:
        logo_fim, dir_ini = partes_bloco(ws, c1, c2, logo_fim)
        mesclar(ws, 1, c1, logo_fim, 2)
        contorno(ws, 1, c1, 2, logo_fim, azul, base)
        logo(ws, c1, logo_fim)
    t1, t2 = logo_fim + 1, dir_ini - 1
    mesclar(ws, 1, t1, t2, 2)
    c = ws.cell(1, t1)
    partes = [TextBlock(InlineFont(rFont=FONTE, sz=10, b=True, color=DESTAQUE), f"{codigo}   ")] if codigo else []
    partes.append(TextBlock(InlineFont(rFont=FONTE_TITULO, sz=16, color="000000"), titulo + ("  (continuação)" if continuacao else "")))
    c.value = CellRichText(*partes)
    c.alignment = CENTRO
    contorno(ws, 1, t1, 2, t2, azul, base)
    if dir_ini <= c2:
        mesclar(ws, 1, dir_ini, c2)
        mesclar(ws, 2, dir_ini, c2)
        r = ws.cell(1, dir_ini, F_REV)
        r.font, r.alignment = fonte(11, "000000", nome_fonte=FONTE_TITULO), Alignment(horizontal="center", vertical="bottom")
        r = ws.cell(2, dir_ini, F_ELAB)
        r.font, r.alignment = fonte(8, CINZA, True, nome_fonte=FONTE_ESTREITA), Alignment(horizontal="center", vertical="top", wrap_text=True)
        contorno(ws, 1, dir_ini, 2, c2, azul, base)
    return t1


def linha_aviso(ws, c1, c2):
    mesclar(ws, 3, c1, c2)
    c = ws.cell(3, c1)
    c.value = CellRichText(TextBlock(InlineFont(rFont=FONTE, sz=7.5, b=True, color=CINZA), "Aviso: "),
                           TextBlock(InlineFont(rFont=FONTE, sz=7.5, i=True, color=CINZA), AVISO))
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = max(ws.row_dimensions[3].height or 0, altura_texto("Aviso: " + AVISO, largura_total(ws, c1, c2) - 2, 7.5, 14))
    ws.row_dimensions[4].height = 4


def topo(ws, codigo, titulo, ncols, faixa_txt=None, c_ini=1):
    """Topo das abas de formulário: bloco, aviso e faixa azul com o nome da seção."""
    t1 = bloco(ws, codigo, titulo, c_ini, ncols)
    linha_aviso(ws, t1, ncols)
    if faixa_txt:
        faixa(ws, LINHA_FAIXA, [(c_ini, ncols, faixa_txt)])
    return t1


def faixa(ws, linha, grupos):
    """Faixa azul do modelo padrão: rótulos em branco, Arial Black."""
    for c1, c2, rot in grupos:
        mesclar(ws, linha, c1, c2)
        c = ws.cell(linha, c1, rot)
        c.font, c.fill, c.alignment = fonte(9, "FFFFFF", nome_fonte=FONTE_TITULO), preench(AZUL), CENTRO
        for i in range(c1, c2 + 1):
            ws.cell(linha, i).fill = preench(AZUL)
        contorno(ws, linha, c1, linha, c2, Side(style="thin", color="FFFFFF"))
    ws.row_dimensions[linha].height = 22


def secao(ws, linha, titulo, c1, c2):
    """Título de seção do modelo padrão (COMO USAR, PARÂMETROS...): Arial Black azul sobre azul-acinzentado."""
    mesclar(ws, linha, c1, c2)
    c = ws.cell(linha, c1, titulo)
    c.font, c.fill, c.alignment = fonte(10, AZUL, nome_fonte=FONTE_TITULO), preench(SECAO_FUNDO), Alignment(vertical="center", indent=1)
    contorno(ws, linha, c1, linha, c2, Side(style="thin", color=SECAO_BORDA))
    ws.row_dimensions[linha].height = 21.75


def cabecalho(ws, linha, titulos_cores, col0=1, espaco=None):
    """Cabeçalho de tabela: azul-claro e texto azul para preencher; mais claro e texto cinza para colunas automáticas."""
    altura = 24
    for i, (titulo, papel) in enumerate(titulos_cores):
        c = ws.cell(linha, col0 + i, titulo)
        calc = papel in ("calculo", "futuro")
        c.font = fonte(8.5, CINZA if calc else AZUL, True)
        c.fill, c.alignment, c.border = preench(CAB_CALC if calc else CAB_FUNDO), CENTRO, B_CAB
        altura = max(altura, altura_texto(titulo, largura(ws, col0 + i) or 9, 8.5, 24))
    ws.row_dimensions[linha].height = min(altura, 40)


# ------------------------------------------------------------------ tabelas
def estilo_dado(cell, col):
    cell.font = fonte(9, CINZA if col.papel in ("calculo", "futuro") else TINTA)
    cell.border = B_GRADE
    if col.tipo in ("sim_nao", "data") or (col.numero or "").startswith('"Sim"') or (col.tipo in ("texto", "lista") and col.largura <= 12):
        cell.alignment = Alignment(horizontal="center", vertical="center")
    else:  # recuo de um caractere: números não encostam no texto da coluna ao lado
        lado = "right" if col.tipo in ("inteiro", "numero_horas") else "left"
        cell.alignment = Alignment(horizontal=lado, vertical="center", indent=1, wrap_text=col.largura >= 28)
    if col.numero:
        cell.number_format = col.numero
    elif col.papel != "calculo" and col.tipo in ("texto", "lista", "sim_nao"):
        cell.number_format = FORMATO_TEXTO
    if col.papel in ("entrada", "manual", "futuro") or (col.papel == "id" and col.id.startswith(("execucao", "falha", "pessoa", "requisito", "ocorrencia", "ativo"))):
        cell.protection = Protection(locked=False)


def escrever_valor(cell, col, valor):
    if valor is None or valor == "":
        return
    if col.tipo == "data":
        if isinstance(valor, str):
            valor = dt.date.fromisoformat(valor)
        cell.value = valor
    elif col.tipo in ("numero_horas", "inteiro", "percentual"):
        cell.value = valor
    else:
        texto(cell, str(valor))


def aplicar_validacoes(ws, colunas, l1, l2, col0=1):
    for i, col in enumerate(colunas):
        if not col.dv:
            continue
        L = letra(col0 + i)
        d = dict(col.dv)
        f1 = d.pop("formula1").replace("{c}", f"{L}{l1}")
        f2 = d.pop("formula2", None)
        erro, prompt = d.pop("error"), d.pop("prompt")
        dv = DataValidation(formula1=f1, formula2=f2, allow_blank=True, showErrorMessage=True, showInputMessage=True,
                            errorTitle="Valor não aceito", error=erro[:255], promptTitle=col.cabecalho[:32], prompt=prompt[:255], **d)
        dv.add(f"{L}{l1}:{L}{l2}")
        ws.add_data_validation(dv)


AMBAR = ("Pendente", "Verificar", "Duplicada", "Tarefa ", "Fora do plano", "Previsão", "Sem vínculo", "Não verificado", "Informe")
COLUNAS_STATUS = ("conferencia", "orientacao", "vinculo", "situacao_plano", "parada_situacao", "tag_conferencia")


def regra_status(ws, alvo, L, l1):
    """Só o que pede ação ganha cor: ERRO em vermelho, conferir em âmbar. OK fica neutro."""
    ws.conditional_formatting.add(alvo, FormulaRule(formula=[f'OR(LEFT({L}{l1},4)="ERRO",LEFT({L}{l1},7)="ATENÇÃO")'], stopIfTrue=True,
                                                    fill=preench(ERRO_FUNDO), font=Font(bold=True, color=SEGURANCA)))
    cond = ",".join(f'LEFT({L}{l1},{len(p)})="{p}"' for p in AMBAR)
    ws.conditional_formatting.add(alvo, FormulaRule(formula=[f"OR({cond})"], stopIfTrue=True, fill=preench(AVISO_FUNDO), font=Font(color=AVISO_TEXTO)))


CLASSES = (("A", "F6DDE1", SEGURANCA), ("B", "FBE8D3", "8A4500"), ("C", AZUL_SUAVE, AZUL), ("Pendente", AVISO_FUNDO, AVISO_TEXTO))


def formatacao_condicional(ws, colunas, l1, l2, col0=1):
    n = len(colunas)
    chave = f"${letra(col0)}{l1}"
    for i, col in enumerate(colunas):
        L = letra(col0 + i)
        alvo = f"{L}{l1}:{L}{l2}"
        if col.id == "criticidade.classe":
            for valor, fundo, cor in CLASSES:
                ws.conditional_formatting.add(alvo, FormulaRule(formula=[f'{L}{l1}="{valor}"'], stopIfTrue=True, fill=preench(fundo),
                                                                font=Font(bold=True, color=cor)))
        elif col.id.rsplit(".", 1)[1] in COLUNAS_STATUS:
            regra_status(ws, alvo, L, l1)
        if col.id == "tarefa.seguranca":
            ws.conditional_formatting.add(alvo, FormulaRule(formula=[f'{L}{l1}="Sim"'], font=Font(bold=True, color=SEGURANCA)))
        if col.id in ("tarefa.ativa", "execucao.cancelada", "pessoa.ativo"):
            alvo_txt = "Não" if col.id != "execucao.cancelada" else "Sim"
            ws.conditional_formatting.add(f"{letra(col0)}{l1}:{letra(col0 + n - 1)}{l2}",
                                          FormulaRule(formula=[f'${L}{l1}="{alvo_txt}"'], font=Font(color="8C96A0", strike=col.id == "execucao.cancelada")))
    ws.conditional_formatting.add(f"{letra(col0)}{l1}:{letra(col0 + n - 1)}{l2}",
                                  FormulaRule(formula=[f'AND({chave}<>"",MOD(ROW(),2)=0)'], fill=preench(ZEBRA)))


def grupos_da_aba(aba):
    """(c1, c2, rótulo) de cada grupo da faixa azul."""
    ids = [c.id for c in aba.colunas]
    defs = E.GRUPOS.get(aba.nome, [(aba.titulo, None)])
    inicios = [1 if cid is None else ids.index(cid) + 1 for _, cid in defs]
    fins = inicios[1:] + [len(ids) + 1]
    return [(a, b - 1, rot.upper()) for (rot, _), a, b in zip(defs, inicios, fins)]


def montar_aba_tabela(wb, aba, linhas, versao_exemplo, arq):
    ws = wb.create_sheet(aba.nome)
    ws.sheet_view.showGridLines = False
    n = len(aba.colunas)
    l1, l2 = E.PRIMEIRA, E.PRIMEIRA + aba.capacidade - 1
    for i, col in enumerate(aba.colunas, start=1):
        ws.column_dimensions[letra(i)].width = col.largura
        if col.oculta:
            ws.column_dimensions[letra(i)].hidden = True
    codigo = aba.codigo or arq.codigo
    revisoes = aba is E.REVISOES_ARQ
    titulo = arq.titulo if revisoes else aba.titulo
    quebra = [c.id for c in aba.colunas].index(E.QUEBRA[aba.nome]) + 1 if aba.nome in E.QUEBRA else None
    rep_fim = ws[aba.colunas_repetidas.split(":")[1] + "1"].column if (quebra and aba.colunas_repetidas) else None
    if quebra:  # duas partes na impressão: cada parte tem seu topo; o logo fica nas colunas repetidas
        t1 = bloco(ws, codigo, titulo, 1, quebra - 1, logo_fim=rep_fim)
        linha_aviso(ws, t1, quebra - 1)
        bloco(ws, codigo, titulo, quebra, n, continuacao=True)
        linha_aviso(ws, quebra, n)
    else:
        t1 = bloco(ws, codigo, titulo, 1, n)
        linha_aviso(ws, t1, n)
    grupos = [(1, n, "HISTÓRICO DE REVISÕES")] if revisoes else grupos_da_aba(aba)
    faixa(ws, LINHA_FAIXA, grupos)
    cabecalho(ws, E.LINHA_CAB, [(c.cabecalho, c.papel) for c in aba.colunas])
    for r in range(l1, l2 + 1):
        dado = linhas[r - l1] if r - l1 < len(linhas) else None
        for i, col in enumerate(aba.colunas):
            cell = ws.cell(r, i + 1)
            estilo_dado(cell, col)
            if col.formula:
                cell.value = col.formula.replace("{r}", str(r))
            elif dado is not None:
                escrever_valor(cell, col, dado.get(col.id))
    for i, col in enumerate(aba.colunas):
        if col.nome:
            nome(wb, col.nome, ref_abs(aba.nome, i + 1, l1, l2))
    aplicar_validacoes(ws, aba.colunas, l1, l2)
    formatacao_condicional(ws, aba.colunas, l1, l2)
    ws.freeze_panes = ws.cell(l1, aba.congelar_col + 1)
    ws.auto_filter.ref = f"A{E.LINHA_CAB}:{letra(n)}{l2}"
    usadas = len(linhas) if versao_exemplo else aba.linhas_impressao_modelo
    configurar_impressao(ws, aba.paisagem, n, l1 + max(usadas, 1) - 1, codigo, titulos=f"1:{E.LINHA_CAB}", quebra=quebra,
                         colunas=aba.colunas_repetidas if quebra else None)
    if aba.protegida:
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.sort = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
    return ws


# ------------------------------------------------------------------ formulários (células únicas)
def campo_entrada(cell, valor=None, fmt=FORMATO_TEXTO):
    """Campo a preencher do modelo padrão: fundo azul-claro, texto azul em negrito."""
    cell.border, cell.font, cell.alignment = B_GRADE, fonte(9, ENTRADA_TEXTO, True), Alignment(vertical="center", wrap_text=True, indent=1)
    cell.fill = preench(ENTRADA_FUNDO)
    cell.number_format = fmt
    cell.protection = Protection(locked=False)
    if valor not in (None, ""):
        if isinstance(valor, str):
            texto(cell, valor)
        else:
            cell.value = valor


def campo_calculo(cell, formula, fmt=None):
    cell.value = formula
    cell.fill, cell.font, cell.border = preench(CALC_FUNDO), fonte(9, CINZA), B_GRADE
    cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    if fmt:
        cell.number_format = fmt


def rotulo(ws, r, c, txt):
    cell = ws.cell(r, c)
    texto(cell, txt)
    cell.font, cell.alignment, cell.border = fonte(9, TINTA), Alignment(vertical="center", wrap_text=True, indent=1), B_GRADE
    cell.fill = preench(CALC_FUNDO)


def nota_curta(ws, r, c, txt):
    cell = ws.cell(r, c)
    texto(cell, txt)
    cell.font, cell.alignment, cell.border = fonte(8.5, CINZA), Alignment(vertical="center", wrap_text=True, indent=1), B_GRADE


def montar_organizacao(wb, dados):
    ws = wb.create_sheet("Organização")
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([25, 27, 26, 32], start=1):
        ws.column_dimensions[letra(i)].width = w
    topo(ws, "PCMD-01", "Organização", 4, "DADOS DA ORGANIZAÇÃO")
    cabecalho(ws, 6, [("Campo", "entrada"), ("Valor", "entrada"), ("Resultado", "calculo"), ("Observação", "entrada")])
    g = dados.get("gestao") if dados else None
    org = g["organizacao"] if g else {"nome": "", "referencial": "GERAL (sem norma)", "padraoCodigo": "PCMD", "elaborador": "", "aprovador": ""}
    pessoa = 'IFERROR(INDEX(cad_pessoa_nome,MATCH(B{r},cad_pessoa_id,0)),"Pessoa não cadastrada")'
    linhas = [
        ("Nome da organização", org["nome"], None, "Ex.: Indústria Exemplo."),
        ("Referencial", org["referencial"], None, "Padrão GERAL (sem norma). Não cite norma sem fonte oficial, versão e data."),
        ("Prefixo dos documentos", org["padraoCodigo"], None, "Padrão PCMD."),
        ("Elaborador padrão (ID)", org["elaborador"], '=IF(B10="","",' + pessoa.format(r=10) + ")", "ID da aba Pessoas."),
        ("Aprovador padrão (ID)", org["aprovador"], '=IF(B11="","A definir (aprovação não registrada)",' + pessoa.format(r=11) + ")",
         "Só informe após aprovação efetiva."),
    ]
    for k, (rot, val, form, obs) in enumerate(linhas):
        r = 7 + k
        rotulo(ws, r, 1, rot)
        campo_entrada(ws.cell(r, 2), val)
        campo_calculo(ws.cell(r, 3), form)
        nota_curta(ws, r, 4, obs)
        ws.row_dimensions[r].height = 22
    rotulo(ws, 12, 1, "Meta de disponibilidade (%)")
    v = g["metas"]["disponibilidade"] if g else None
    campo_entrada(ws.cell(12, 2), None if v is None else v / 100, "0.0%")
    campo_calculo(ws.cell(12, 3), '=IF(B12="","A definir",B12)', "0.0%")
    nota_curta(ws, 12, 4, "Disponibilidade, MTBF e MTTR não são calculados: regras técnicas pendentes.")
    ws.row_dimensions[12].height = 22
    for nm, cel in (("cad_org_nome", "B7"), ("cad_referencial", "B8"), ("cad_prefixo", "B9"), ("cad_meta_disponibilidade", "B12")):
        nome(wb, nm, f"'Organização'!${cel[0]}${cel[1:]}")
    for celula in ("B10", "B11"):
        dv = DataValidation(type="list", formula1="=cad_pessoa_id", allow_blank=True, showErrorMessage=True,
                            error="Use um ID da aba Pessoas.", errorTitle="Valor não aceito")
        dv.add(celula)
        ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True, showErrorMessage=True,
                        errorTitle="Meta não aceita", error="Informe um percentual entre 0% e 100%. Vazio = A definir.")
    dv.add("B12")
    ws.add_data_validation(dv)
    configurar_impressao(ws, False, 4, 12, "PCMD-01")


def montar_criterios(wb, dados):
    ws = wb.create_sheet("Critérios")
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([26, 22, 22, 22, 26], start=1):
        ws.column_dimensions[letra(i)].width = w
    topo(ws, "PCMD-02", "Critérios de criticidade", 5, "LIMITES E NÍVEIS DE CRITICIDADE")
    cabecalho(ws, 6, [("Campo", "entrada"), ("Valor", "entrada"), ("Resultado", "calculo"), ("Regra", "entrada")])
    g = dados.get("gestao") if dados else None
    lim = g["criticidade"]["limites"] if g else {"a": 12, "b": 9}
    for r, rot, val, regra in ((7, "Limite A (soma mínima)", lim["a"], "Inteiros com 5 ≤ B < A ≤ 15."), (8, "Limite B (soma mínima)", lim["b"], "Padrão inicial: A = 12 e B = 9.")):
        rotulo(ws, r, 1, rot)
        campo_entrada(ws.cell(r, 2), val, "0")
        nota_curta(ws, r, 4, regra)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.cell(r, 3).border = B_GRADE
        ws.row_dimensions[r].height = 22
    rotulo(ws, 9, 1, "Limites válidos?")
    campo_calculo(ws.cell(9, 2), "=IF(AND(ISNUMBER(B7),ISNUMBER(B8)),AND(B7=INT(B7),B8=INT(B8),B8>=5,B8<B7,B7<=15),FALSE)")
    campo_calculo(ws.cell(9, 3), '=IF(B9,"OK","ERRO: Inválidos; use inteiros com 5 ≤ B < A ≤ 15 (classes ficam Pendente)")')
    ws.merge_cells("C9:E9")
    ws.row_dimensions[9].height = 22
    regra_status(ws, "C9", "C", 9)
    nome(wb, "cad_limite_a", "'Critérios'!$B$7")
    nome(wb, "cad_limite_b", "'Critérios'!$B$8")
    nome(wb, "cad_limites_ok", "'Critérios'!$B$9")
    dv = DataValidation(type="whole", operator="between", formula1="5", formula2="15", allow_blank=False, showErrorMessage=True,
                        errorTitle="Limite não aceito", error="Use inteiro entre 5 e 15, com B menor que A.")
    dv.add("B7:B8")
    ws.add_data_validation(dv)
    secao(ws, 11, "Significado dos níveis (definido pelo mantenedor)", 1, 5)
    cabecalho(ws, 12, [("Critério", "entrada"), ("Nível 1", "manual"), ("Nível 2", "manual"), ("Nível 3", "manual"), ("Situação", "calculo")], espaco=False)
    for k, crit in enumerate(["N1 Segurança", "N2 Qualidade", "N3 Produção", "N4 Frequência de falha", "N5 Custo/tempo de reparo"]):
        r = 13 + k
        rotulo(ws, r, 1, crit)
        for c in (2, 3, 4):
            campo_entrada(ws.cell(r, c))
        campo_calculo(ws.cell(r, 5), f'=IF(COUNTA(B{r}:D{r})<3,"Pendente: a definir pelo mantenedor","OK")')
        ws.row_dimensions[r].height = 30
    regra_status(ws, "E13:E17", "E", 13)
    rotulo(ws, 19, 1, "Regra")
    mesclar(ws, 19, 2, 5)
    nota_curta(ws, 19, 2, "Segurança = 3 → A; senão soma ≥ A → A; soma ≥ B → B; senão C. Nota ausente ou limite inválido → Pendente. "
                          "A classe não altera frequências.")
    ws.row_dimensions[19].height = 30
    configurar_impressao(ws, False, 5, 19, "PCMD-02")


# ------------------------------------------------------------------ listas
def especialidades(biblioteca):
    vistas = []
    for t in biblioteca["tipos"]:
        for tarefa in t["tarefas"]:
            if tarefa["especialidade"] not in vistas:
                vistas.append(tarefa["especialidade"])
    return vistas


def blocos_listas(biblioteca):
    return {
        "frequencias": (["Frequência", "Execuções em 52 semanas", "Observação"], [list(f) for f in E.FREQUENCIAS],
                        [("lst_frequencias", 0, 8), ("lst_frequencias_vezes", 1, 8)], [12, 12, 44]),
        "especialidades": (["Especialidade"], [[e] for e in especialidades(biblioteca)], [("lst_especialidades", 0, 30)], [16]),
        "tipos": (["Tipo (ID)", "Tipo (nome)", "Sigla"], [[t["id"], t["nome"], t["sigla"]] for t in biblioteca["tipos"]],
                  [("lst_tipo_id", 0, 30), ("lst_tipo_nome", 1, 30), ("lst_tipo_sigla", 2, 30)], [20, 22, 8]),
        "sim_nao": (["Sim/Não"], [["Sim"], ["Não"]], [("lst_sim_nao", 0, 2)], [9]),
        "origem_execucao": (["Origem da execução"], [["app"], ["legado"]], [("lst_origem_execucao", 0, 2)], [12]),
        "status_requisito": (["Status de requisito"], [["Não verificado"], ["Verificado"]], [("lst_status_requisito", 0, 2)], [15]),
        "situacao_documento": (["Situação do documento"], [[s] for s in E.SITUACOES_DOCUMENTO], [("lst_situacao_documento", 0, 10)], [16]),
        "tipo_documento": (["Tipo de documento"], [[t] for t in E.TIPOS_DOCUMENTO], [("lst_tipo_documento", 0, 10)], [15]),
        "tipos_tarefa": (["Tipos de tarefa (consulta)"], [[t] for t in E.TIPOS_TAREFA_REFERENCIA], [], [18]),
    }


def montar_listas(wb, biblioteca, chaves, codigo):
    ws = wb.create_sheet("Listas")
    ws.sheet_view.showGridLines = False
    todos = blocos_listas(biblioteca)
    col0, ultima, preenchidas = 1, 1, 1
    posicoes = []
    for k in chaves:
        titulos, valores, nomes, larguras = todos[k]
        for i, w in enumerate(larguras):
            ws.column_dimensions[letra(col0 + i)].width = w
        posicoes.append((col0, k))
        ultima = col0 + len(titulos) - 1
        col0 = ultima + 2
        if col0 - 1 > ultima:
            ws.column_dimensions[letra(col0 - 1)].width = 2
    ncols = ultima
    falta = 86 - largura_total(ws, 1, ultima)
    if falta > 0:  # poucas listas: uma coluna vazia ao fim dá largura ao bloco de cabeçalho
        ncols = ultima + 2 if ultima + 1 == col0 - 1 else ultima + 1
        ws.column_dimensions[letra(ncols)].width = max(falta - (2 if ncols > ultima + 1 else 0), 8)
    topo(ws, codigo, "Listas", ncols, "VALORES DAS LISTAS SUSPENSAS")
    for col0, k in posicoes:
        titulos, valores, nomes, larguras = todos[k]
        cabecalho(ws, E.LINHA_CAB, [(t, "referencia") for t in titulos], col0=col0)
        preenchidas = max(preenchidas, len(valores))
        for kk, linha in enumerate(valores):  # os intervalos nomeados vão além: sobram linhas para novos itens
            for i, v in enumerate(linha):
                cell = ws.cell(E.PRIMEIRA + kk, col0 + i)
                cell.border, cell.font = B_GRADE, fonte(9)
                cell.alignment = Alignment(vertical="center", indent=1)
                if isinstance(v, str):
                    texto(cell, v)
                else:
                    cell.value = v
        for nm, desloc, tam in nomes:
            nome(wb, nm, ref_abs("Listas", col0 + desloc, E.PRIMEIRA, E.PRIMEIRA + tam - 1))
    ws.freeze_panes = "A7"
    configurar_impressao(ws, ncols > 8, ncols, E.PRIMEIRA + preenchidas - 1, codigo, titulos=f"1:{E.LINHA_CAB}")


# ------------------------------------------------------------------ cronograma (painel)
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]


def montar_cronograma(wb, dados):
    ws = wb.create_sheet("Cronograma")
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([6, 26, 15, 15, 18, 15, 6], start=1):
        ws.column_dimensions[letra(i)].width = w
    ws.column_dimensions["A"].hidden = True  # número do mês, usado só nas fórmulas
    ws.column_dimensions["G"].hidden = True  # mês abreviado, usado só no eixo do gráfico
    topo(ws, "PCMD-04", "Cronograma", 7, "RESUMO DO ANO  ·  MÊS DA DATA PREVISTA  ·  % = CONCLUÍDAS ÷ PREVISTAS")
    secao(ws, 6, "PARÂMETROS", 2, 6)
    meta = (dados["gestao"]["metas"]["aderencia"] if dados else None)
    par = [
        (7, "Ano consultado", dados["ano"] if dados else None, "0", True),
        (8, "Início do ciclo", dt.date.fromisoformat(dados["inicio"]) if dados else None, "dd/mm/yyyy", True),
        (9, "Fonte das ocorrências", "PCM Direto: 52 semanas desde 03/11/2025 (fictício)" if dados else None, FORMATO_TEXTO, True),
        (10, "Conferência do ano", '=IF(ISNUMBER(crono_ano),IF(AND(crono_ano=INT(crono_ano),crono_ano>=1900,crono_ano<=9999),"OK","ERRO: ano inválido"),"Informe o ano consultado")', None, False),
        (11, "Meta de aderência", None if meta is None else meta / 100, "0.0%", True),
    ]
    for r, rot, val, fmt, ent in par:
        rotulo(ws, r, 2, rot)
        fim = 6 if r == 9 else 4
        mesclar(ws, r, 3, fim)
        c = ws.cell(r, 3)
        if ent:
            campo_entrada(c, val, fmt)
            for i in range(4, fim + 1):
                ws.cell(r, i).border = B_CAIXA
        else:
            campo_calculo(c, val, fmt)
        ws.row_dimensions[r].height = 20
        ws.cell(r, 3).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c = ws.cell(11, 5, '=IF(crono_meta_aderencia="","A definir",IF(NOT(ISNUMBER(E27)),"",IF(E27>=crono_meta_aderencia,"Na meta","Abaixo da meta")))')
    c.font, c.alignment = fonte(9.5, CINZA, italico=True), Alignment(vertical="center", indent=1)
    regra_status(ws, "C10", "C", 10)
    nome(wb, "crono_ano", "'Cronograma'!$C$7")
    nome(wb, "crono_inicio", "'Cronograma'!$C$8")
    nome(wb, "crono_meta_aderencia", "'Cronograma'!$C$11")
    for celula, tipo, f1, f2, erro in (("C7", "whole", "1900", "9999", "Informe o ano com quatro dígitos."),
                                       ("C8", "date", "1", "2958465", "Informe uma data válida."),
                                       ("C11", "decimal", "0", "1", "Informe um percentual entre 0% e 100%. Vazio = A definir.")):
        dv = DataValidation(type=tipo, operator="between", formula1=f1, formula2=f2, allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor não aceito", error=erro)
        dv.add(celula)
        ws.add_data_validation(dv)

    cr = E.CRONO
    j, d, t = cr["jan"], cr["jan"] + 11, cr["total"]
    assert t == 27  # a situação frente à meta (E11) lê o total em E27

    def calc(r, c, formula, fmt=None, negrito=False):
        cell = ws.cell(r, c, formula)
        cell.font, cell.border = fonte(9, TINTA if c <= 2 else CINZA, negrito), B_GRADE
        cell.alignment = Alignment(horizontal="right" if c >= 3 else None, vertical="center", wrap_text=c == 2, indent=1 if c >= 3 else 0)
        if fmt:
            cell.number_format = fmt
        return cell

    secao(ws, cr["cab_mes"] - 1, "RESUMO MENSAL", 2, 6)
    cabecalho(ws, cr["cab_mes"], [(x, "calculo") for x in ("Mês (nº)", "Mês", "Previstas", "Concluídas", "% realizado", "Horas estimadas")], espaco=False)
    for k, m in enumerate(MESES):
        r = j + k
        calc(r, 1, k + 1, "0")
        ws.cell(r, 7, m[:3])
        calc(r, 2, m)
        calc(r, 3, f'=IF(NOT(ISNUMBER(crono_ano)),"",SUMIFS(oc_conta,oc_ano,crono_ano,oc_mes,A{r}))', "0")
        calc(r, 4, f'=IF(C{r}="","",SUMIFS(oc_concluida,oc_ano,crono_ano,oc_mes,A{r}))', "0")
        calc(r, 5, f'=IF(C{r}="","",IF(C{r}=0,"Sem previsão",D{r}/C{r}))', "0.0%")
        calc(r, 6, f'=IF(C{r}="","",SUMIFS(oc_horas,oc_conta,1,oc_ano,crono_ano,oc_mes,A{r}))', "0.00")
    ws.conditional_formatting.add(f"B{j}:F{d}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=preench(ZEBRA)))
    calc(t, 1, "")
    calc(t, 2, "Total do ano", negrito=True)
    calc(t, 3, f'=IF(NOT(ISNUMBER(crono_ano)),"",SUM(C{j}:C{d}))', "0", True)
    calc(t, 4, f'=IF(C{t}="","",SUM(D{j}:D{d}))', "0", True)
    calc(t, 5, f'=IF(C{t}="","",IF(C{t}=0,"Sem previsão",D{t}/C{t}))', "0.0%", True)
    calc(t, 6, f'=IF(C{t}="","",SUM(F{j}:F{d}))', "0.00", True)
    for c in range(2, 7):
        ws.cell(t, c).font = fonte(9, AZUL, True)
        ws.cell(t, c).fill = preench(CAB_FUNDO)
        ws.cell(t, c).border = Border(left=FINO, right=FINO, top=Side(style="medium", color=AZUL), bottom=Side(style="medium", color=AZUL))
    ws.conditional_formatting.add(f"E{j}:E{d}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="8FAECB", showValue=True))
    ws.conditional_formatting.add(f"E{j}:E{d}", FormulaRule(formula=[f'AND(ISNUMBER(E{j}),ISNUMBER(crono_meta_aderencia),E{j}<crono_meta_aderencia)'],
                                                          font=Font(bold=True, color=AVISO_TEXTO)))

    graf = BarChart()
    graf.type = "col"
    graf.grouping = "clustered"
    graf.overlap = -10
    graf.gapWidth = 70
    graf.y_axis.majorGridlines = None
    graf.y_axis.delete = False
    graf.x_axis.delete = False
    graf.legend.position = "b"
    graf.add_data(Reference(ws, min_col=3, max_col=4, min_row=cr["cab_mes"], max_row=d), titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=7, min_row=j, max_row=d))
    graf.visible_cells_only = False  # o eixo usa a coluna G, oculta
    for serie, cor in zip(graf.series, (AZUL, "9FB4C7")):
        serie.graphicalProperties.solidFill = cor
        serie.graphicalProperties.line.solidFill = cor
    graf.height, graf.width = 6.6, 16.5
    secao(ws, cr["grafico"] - 1, "PREVISTAS × CONCLUÍDAS", 2, 6)
    ws.add_chart(graf, f"B{cr['grafico']}")

    f0 = cr["fora_cab"]
    secao(ws, f0 - 1, "FORA DO ANO CONSULTADO", 2, 6)
    cabecalho(ws, f0, [(x, "calculo") for x in ("", "Posição", "Previstas", "Concluídas", "% realizado", "Horas estimadas")], espaco=False)
    for k, pos in enumerate(["Antes do ano", "Depois do ano"]):
        r = f0 + 1 + k
        calc(r, 1, "")
        calc(r, 2, pos)
        calc(r, 3, f'=IF(NOT(ISNUMBER(crono_ano)),"",SUMIFS(oc_conta,oc_posicao,B{r}))', "0")
        calc(r, 4, f'=IF(C{r}="","",SUMIFS(oc_concluida,oc_posicao,B{r}))', "0")
        calc(r, 5, f'=IF(C{r}="","",IF(C{r}=0,"Sem previsão",D{r}/C{r}))', "0.0%")
        calc(r, 6, f'=IF(C{r}="","",SUMIFS(oc_horas,oc_conta,1,oc_posicao,B{r}))', "0.00")

    def quadro(c0, titulo, cab, linhas):
        secao(ws, c0 - 1, titulo.upper(), 2, 6)
        cabecalho(ws, c0, [(x, "calculo") for x in ("", cab[0], cab[1], cab[2], "", "")], espaco=False)
        mesclar(ws, c0, 4, 6)
        for k, (rot, form, obs) in enumerate(linhas):
            r = c0 + 1 + k
            calc(r, 1, "")
            calc(r, 2, rot)
            calc(r, 3, form, "0").alignment = Alignment(horizontal="center", vertical="center")
            mesclar(ws, r, 4, 6)
            calc(r, 4, obs).alignment = Alignment(vertical="center", indent=1)
            for c in (5, 6):
                ws.cell(r, c).border = B_GRADE

    quadro(cr["semv_cab"], "Execuções fora do percentual (ano da data realizada)", ("Situação", "Registros", "Observação"), [
        ("Sem vínculo (legado ou sem data prevista)", '=IF(NOT(ISNUMBER(crono_ano)),"",COUNTIFS(ex_vinculo,"Sem vínculo",ex_cancelada,"<>Sim",ex_ano,crono_ano))', "Não contam."),
        ("Data prevista sem ocorrência", '=IF(NOT(ISNUMBER(crono_ano)),"",COUNTIFS(ex_vinculo,"Previsão não encontrada",ex_cancelada,"<>Sim",ex_ano,crono_ano))', "Conferir a data prevista."),
        ("Canceladas", '=IF(NOT(ISNUMBER(crono_ano)),"",COUNTIFS(ex_cancelada,"Sim",ex_ano,crono_ano))', "Preservadas; não contam."),
    ])
    l2o = E.PRIMEIRA + E.OCORRENCIAS.capacidade - 1
    l2x = E.PRIMEIRA + E.EXECUCOES.capacidade - 1
    c0 = cr["conf_cab"]
    quadro(c0, "Conferências", ("Verificação", "Resultado", "Como corrigir"), [
        ("Ocorrências com ERRO", '=COUNTIF(oc_conferencia,"ERRO*")', "Coluna Conferência em Ocorrências."),
        ("Ocorrências duplicadas", '=COUNTIF(oc_conferencia,"Duplicada*")', "Remova a linha repetida."),
        ("Capacidade de Ocorrências", f"=IF(COUNTA('Ocorrências'!A{l2o + 1}:A{l2o + 5000})>0,\"ATENÇÃO: há linhas fora das fórmulas\",\"OK\")",
         "Estenda fórmulas e nomes (COMO-USAR)."),
        ("Capacidade de Execuções", f"=IF(COUNTA('Execuções'!A{l2x + 1}:A{l2x + 5000})>0,\"ATENÇÃO: há linhas fora das fórmulas\",\"OK\")",
         "Estenda fórmulas e nomes (COMO-USAR)."),
    ])
    ws.conditional_formatting.add(f"C{c0 + 1}:C{c0 + 4}", FormulaRule(formula=[f'LEFT(C{c0 + 1},8)="ATENÇÃO"'], fill=preench(ERRO_FUNDO),
                                                                     font=Font(bold=True, color=SEGURANCA)))
    ws.conditional_formatting.add(f"C{c0 + 1}:C{c0 + 2}", FormulaRule(formula=[f"N(C{c0 + 1})>0"], fill=preench(AVISO_FUNDO), font=Font(color=AVISO_TEXTO)))
    configurar_impressao(ws, False, 7, c0 + 4, "PCMD-04", titulos="1:4")
    ws.row_breaks.append(Break(id=f0 - 2))


# ------------------------------------------------------------------ config (padrão PCM Direto)
def montar_config(wb, arq, dados):
    """Aba Config do modelo padrão: como usar, parâmetros do documento (lidos pelo cabeçalho de todas as abas) e conferências."""
    ws = wb.create_sheet("Config")
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([20, 16, 26, 20, 16, 16], start=1):
        ws.column_dimensions[letra(i)].width = w
    topo(ws, arq.codigo, arq.titulo, 6, "CONFIGURAÇÕES")
    r = 7
    secao(ws, r, "COMO USAR", 1, 6)
    linhas = [f"{k})  {p}" for k, p in enumerate(arq.passos, start=1)]
    linhas.append("Campos em azul: preencher. Texto cinza: automático (fórmula). ERRO em vermelho e Pendente em amarelo pedem ação.")
    for txt in linhas:
        r += 1
        mesclar(ws, r, 1, 6)
        c = ws.cell(r, 1)
        texto(c, txt)
        c.font, c.alignment = fonte(9), Alignment(vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[r].height = altura_texto(txt, 110, 9, 19)
        contorno(ws, r, 1, r, 6, FINO)
    r += 1
    for aba, uso in arq.abas:
        if aba == "Config":
            continue
        c = ws.cell(r, 1, aba)
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{aba}'!A1", display=aba)  # vínculo interno, sem relação externa
        c.font, c.alignment, c.border = fonte(9, AZUL, True, sublinhado="single"), Alignment(vertical="center", indent=1), B_GRADE
        mesclar(ws, r, 2, 6)
        c = ws.cell(r, 2)
        texto(c, uso)
        c.font, c.alignment = fonte(9), Alignment(vertical="center", wrap_text=True, indent=1)
        contorno(ws, r, 2, r, 6, FINO)
        ws.row_dimensions[r].height = altura_texto(uso, 90, 9, 19)
        r += 1

    r += 1
    secao(ws, r, "PARÂMETROS DO DOCUMENTO", 1, 4)
    r += 1
    for c1, c2, rot in ((1, 2, "Parâmetro"), (3, 4, "Valor")):
        mesclar(ws, r, c1, c2)
        cabecalho(ws, r, [(rot, "entrada")], col0=c1)
        ws.cell(r, c2).border = B_CAB
    exemplo = bool(dados)
    if arq.chave == "01":  # a revisão do próprio PCMD-01 está na Lista Mestra
        lm = E.LISTA_MESTRA
        l1, l2 = E.PRIMEIRA, E.PRIMEIRA + lm.capacidade - 1
        ids = [c.id for c in lm.colunas]
        cod, rev, dat = (letra(ids.index(x) + 1) for x in ("documento.codigo", "documento.revisao", "documento.data"))
        alvo = f"MATCH(\"PCMD-01\",'Lista Mestra'!${cod}${l1}:${cod}${l2},0)"
        f_rev = f"=IFERROR(INDEX('Lista Mestra'!${rev}${l1}:${rev}${l2},{alvo})&\"\",\"\")"
        f_dat = f"=IFERROR(IF(INDEX('Lista Mestra'!${dat}${l1}:${dat}${l2},{alvo})=\"\",\"\",INDEX('Lista Mestra'!${dat}${l1}:${dat}${l2},{alvo})),\"\")"
        orig_rev = "automático: Lista Mestra, linha PCMD-01"
    else:
        f_rev = '=IF(COUNTA(rv_revisao)=0,"",INDEX(rv_revisao,COUNTA(rv_revisao))&"")'
        f_dat = '=IF(COUNTA(rv_revisao)=0,"",IF(INDEX(rv_data,COUNTA(rv_revisao))="","",INDEX(rv_data,COUNTA(rv_revisao))))'
        orig_rev = "automático: última linha de Revisões"
    params = [
        ("Elaborado por", "cfg_elaborado", "Pessoa A" if exemplo else None, FORMATO_TEXTO, True),
        ("Aprovado por", "cfg_aprovado", None, FORMATO_TEXTO, True),
        (f"Revisão atual ({orig_rev})", "cfg_revisao", f_rev, FORMATO_TEXTO, False),
        ("Data da revisão (automático)", "cfg_data_revisao", f_dat, "dd/mm/yyyy", False),
        ("Origem dos dados", "pcmd_origem", "Demonstração fictícia gerada com a lógica do PCM Direto (não é exportação do app)" if exemplo else None,
         FORMATO_TEXTO, True),
        ("Data de emissão", "pcmd_emissao", dt.date(2026, 10, 23) if exemplo else None, "dd/mm/yyyy", True),
    ]
    for rot, nm, val, fmt, ent in params:
        r += 1
        mesclar(ws, r, 1, 2)
        rotulo(ws, r, 1, rot)
        ws.cell(r, 2).border = B_GRADE
        mesclar(ws, r, 3, 4)
        if ent:
            campo_entrada(ws.cell(r, 3), val, fmt)
            ws.cell(r, 3).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        else:
            campo_calculo(ws.cell(r, 3), val, None if fmt == FORMATO_TEXTO else fmt)
            ws.cell(r, 3).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.cell(r, 4).border = B_GRADE
        ws.row_dimensions[r].height = 28 if len(rot) > 30 else 19.5
        nome(wb, nm, f"'Config'!$C${r}")
        if fmt == "dd/mm/yyyy" and ent:
            dv = DataValidation(type="date", operator="between", formula1="1", formula2="2958465", allow_blank=True, showErrorMessage=True,
                                errorTitle="Data não aceita", error="Informe uma data válida (dd/mm/aaaa).")
            dv.add(f"C{r}")
            ws.add_data_validation(dv)
        if nm == "cfg_aprovado":
            nota = ws.cell(r, 5, "Só após aprovação efetiva.")
            nota.font, nota.alignment = fonte(8.5, CINZA, italico=True), Alignment(vertical="center", indent=1)

    r += 2
    secao(ws, r, "CONFERÊNCIAS", 1, 6)
    r += 1
    mesclar(ws, r, 4, 6)
    cabecalho(ws, r, [("Aba", "calculo"), ("Linhas preenchidas", "calculo"), ("Linhas com ERRO", "calculo"), ("Capacidade", "calculo")])
    for c in (5, 6):
        ws.cell(r, c).border = B_CAB
    r0 = r + 1
    for aba in arq.tabelas:
        if aba in (E.REVISOES, E.REVISOES_ARQ):
            continue
        r += 1
        l1, l2 = E.PRIMEIRA, E.PRIMEIRA + aba.capacidade - 1
        ids = [c.id for c in aba.colunas]
        conf = next((i + 1 for i, x in enumerate(ids) if ".conferencia" in x), None)
        q = f"'{aba.nome}'!"
        celulas = [(1, aba.nome, None), (2, f"=COUNTA({q}$A${l1}:$A${l2})", "0"),
                   (3, f'=COUNTIF({q}${letra(conf)}${l1}:${letra(conf)}${l2},"ERRO*")' if conf else "—", "0"),
                   (4, f'=IF(COUNTA({q}$A${l2 + 1}:$A${l2 + 2000})>0,"ATENÇÃO: há dados além da linha {l2}; estenda fórmulas e nomes","OK")', None)]
        for c, v, fmt in celulas:
            cell = ws.cell(r, c, v)
            cell.font, cell.border = fonte(9, TINTA if c == 1 else CINZA), B_GRADE
            cell.alignment = Alignment(horizontal="center" if c in (2, 3) else "left", vertical="center", indent=0 if c in (2, 3) else 1, wrap_text=c == 4)
            if fmt:
                cell.number_format = fmt
        mesclar(ws, r, 4, 6)
        for c in (5, 6):
            ws.cell(r, c).border = B_GRADE
        ws.row_dimensions[r].height = 19.5
    regra_status(ws, f"D{r0}:D{r}", "D", r0)
    ws.conditional_formatting.add(f"C{r0}:C{r}", FormulaRule(formula=[f"N(C{r0})>0"], fill=preench(ERRO_FUNDO), font=Font(bold=True, color=SEGURANCA)))
    r += 2
    mesclar(ws, r, 1, 6)
    c = ws.cell(r, 1)
    texto(c, f"Modelo {VERSAO_MODELO} · {AUTOR} · {CONTATO} · Licença {LICENCA} · Sem macros, vínculos externos ou envio de dados. "
             "Regras completas: COMO-USAR.md.")
    c.font, c.alignment = fonte(8, "8C96A0"), Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 24
    configurar_impressao(ws, False, 6, r, arq.codigo)


# ------------------------------------------------------------------ montagem
def gerar(arq, destino, dados=None):
    with open(RAIZ_REPO / "src" / "dados" / "biblioteca.json", encoding="utf-8") as f:
        biblioteca = json.load(f)
    wb = Workbook()
    wb.remove(wb.active)
    linhas = linhas_dados(dados) if dados else {}
    rev_arq = linhas.get("revisoes_arquivo", [])
    if arq.chave == "01":
        montar_aba_tabela(wb, E.LISTA_MESTRA, linhas.get("lista", [{"documento.codigo": c, "documento.titulo": t, "documento.tipo": tp}
                                                                  for c, t, tp in E.DOCUMENTOS_PCMD]), True, arq)
        montar_aba_tabela(wb, E.REVISOES, linhas.get("revisoes", []), bool(dados), arq)
        montar_organizacao(wb, dados)
        montar_aba_tabela(wb, E.PESSOAS, linhas.get("pessoas", []), bool(dados), arq)
        montar_aba_tabela(wb, E.REQUISITOS, linhas.get("requisitos", []), bool(dados), arq)
    elif arq.chave == "02":
        montar_aba_tabela(wb, E.ATIVOS, linhas.get("ativos", []), bool(dados), arq)
        montar_criterios(wb, dados)
        montar_aba_tabela(wb, E.SETORES, linhas.get("setores", []), bool(dados), arq)
    elif arq.chave == "03":
        montar_aba_tabela(wb, E.PLANO, linhas.get("plano", []), bool(dados), arq)
        montar_cronograma(wb, dados)
        montar_aba_tabela(wb, E.OCORRENCIAS, linhas.get("ocorrencias", []), bool(dados), arq)
        montar_aba_tabela(wb, E.EXECUCOES, linhas.get("execucoes", []), bool(dados), arq)
    elif arq.chave == "05":
        montar_aba_tabela(wb, E.FALHAS, linhas.get("falhas", []), bool(dados), arq)
    if E.REVISOES_ARQ in arq.tabelas:
        montar_aba_tabela(wb, E.REVISOES_ARQ, rev_arq, bool(dados), arq)
    montar_config(wb, arq, dados)
    if arq.listas:
        montar_listas(wb, biblioteca, arq.listas, arq.codigo)
    ordem = [a for a, _ in arq.abas]
    assert sorted(ordem) == sorted(wb.sheetnames), (ordem, wb.sheetnames)
    wb._sheets = [wb[n] for n in ordem]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode = "auto"
    wb.properties.creator = "PCM Direto (gerador de modelos)"
    wb.properties.title = f"PCM Direto · {arq.codigo} {arq.titulo}"
    wb.properties.subject = AVISO
    wb.properties.keywords = f"PCM Direto; modelo {VERSAO_MODELO}; licença MIT"
    wb.properties.created = wb.properties.modified = dt.datetime(2026, 10, 23, 12, 0, 0)
    wb.save(destino)
    corrigir_quebra_rodape(destino)


def corrigir_quebra_rodape(caminho):
    """O openpyxl grava a quebra de linha do rodapé como _x000a_; gravamos o caractere de nova linha real."""
    import shutil
    import zipfile
    temp = str(caminho) + ".tmp"
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                texto_xml = dados.decode("utf-8")
                if "<oddFooter>" in texto_xml:
                    ini = texto_xml.index("<oddFooter>")
                    fim = texto_xml.index("</oddFooter>")
                    texto_xml = texto_xml[:ini] + texto_xml[ini:fim].replace("_x000a_", "\n") + texto_xml[fim:]
                    dados = texto_xml.encode("utf-8")
            zout.writestr(item, dados)
    shutil.move(temp, caminho)


def linhas_dados(d):
    g = d["gestao"]
    pessoas = {p["id"]: p["nome"] for p in g["pessoas"]}
    ativos = []
    for e in d["equipamentos"]:
        notas = g["criticidade"]["notas"].get(e["id"], [None] * 5)
        ativos.append({"ativo.id": e["id"], "ativo.tag": e["tag"], "ativo.descricao": e["descricao"], "ativo.tipo_id": e["tipo"],
                       "ativo.fabricante": e.get("fabricante", ""), "ativo.modelo": e.get("modelo", ""), "ativo.setor": e["setor"],
                       "ativo.numero": e["numero"], "criticidade.nota_seguranca": notas[0], "criticidade.nota_qualidade": notas[1],
                       "criticidade.nota_producao": notas[2], "criticidade.nota_frequencia_falha": notas[3], "criticidade.nota_custo_reparo": notas[4]})
    plano = []
    for l in sorted(d["plano"], key=lambda x: (x["tag"], x["tarefaId"])):
        resp = g["responsaveis"].get(f"{l['equipamentoId']}|{l['tarefaId']}", "")
        plano.append({"tarefa.ativo_id": l["equipamentoId"], "tarefa.id": l["tarefaId"], "tarefa.tag": l["tag"], "tarefa.atividade": l["descricao"],
                      "tarefa.origem": l["origem"], "tarefa.versao": l["versao"], "tarefa.frequencia_referencia": l["periodicidadeOriginal"],
                      "tarefa.seguranca": "Sim" if l["seguranca"] else "Não", "tarefa.frequencia_vigente": l["periodicidade"],
                      "tarefa.especialidade": l["especialidade"], "tarefa.duracao_horas": l["duracaoHoras"], "tarefa.ativa": "Sim" if l["ativa"] else "Não",
                      "tarefa.responsavel_id": resp, "tarefa.responsavel_nome": pessoas.get(resp, ""), "tarefa.observacao": l["observacao"]})
    fonte_oc = "Cronograma PCM Direto"
    ocorrencias = [{"ocorrencia.ativo_id": o["equipamentoId"], "ocorrencia.tarefa_id": o["tarefaId"], "ocorrencia.data_prevista": o["dataPrevista"],
                    "ocorrencia.fonte": fonte_oc} for o in d["ocorrencias"]]
    execucoes = [{"execucao.id": h["id"], "execucao.ativo_id": h["equipamentoId"], "execucao.tarefa_id": h["tarefaId"], "execucao.tag": h["tag"],
                  "execucao.descricao": h["descricao"], "execucao.data_prevista": h["dataPrevista"], "execucao.data": h["data"],
                  "execucao.observacao": h["observacao"], "execucao.responsavel_id": h["responsavelId"],
                  "execucao.responsavel_nome": pessoas.get(h["responsavelId"], ""), "execucao.origem": h["origem"],
                  "execucao.cancelada": "Sim" if h["cancelada"] else "Não"} for h in g["historico"]]
    falhas = [{"falha.id": f["id"], "falha.ativo_id": f["equipamentoId"], "falha.data": f["data"], "falha.tag": f["tag"], "falha.sintoma": f["sintoma"],
               "falha.causa": f["causa"], "falha.acao": f["acao"], "falha.horas_parada": f["horasParada"], "falha.peca": f["peca"],
               "falha.responsavel_id": f["responsavelId"], "falha.responsavel_nome": pessoas.get(f["responsavelId"], "")} for f in g["falhas"]]
    lista = [{"documento.codigo": c, "documento.titulo": t, "documento.tipo": tp, "documento.revisao": "00", "documento.data": "2026-10-23",
              "documento.elaborador": "Pessoa A", "documento.situacao": "Em elaboração", "documento.observacao": "Exemplo fictício"}
             for c, t, tp in E.DOCUMENTOS_PCMD]
    revisoes = [{"revisao.documento": c, "revisao.numero": "00", "revisao.data": "2026-10-23", "revisao.descricao": "Emissão inicial (exemplo fictício)",
                 "revisao.responsavel": "Pessoa A"} for c, _, _ in E.DOCUMENTOS_PCMD]
    setores = [{"setor.sigla": s, "setor.descricao": ""} for s in g["organizacao"]["setores"]]
    pessoas_l = [{"pessoa.id": p["id"], "pessoa.nome": p["nome"], "pessoa.funcao": p["funcao"], "pessoa.especialidade": p["especialidade"],
                  "pessoa.ativo": "Sim" if p["ativo"] else "Não"} for p in g["pessoas"]]
    reqs = [{"requisito.id": q["id"], "requisito.referencia": q["referencia"], "requisito.versao": q["versao"], "requisito.data": q["data"],
             "requisito.fonte": q["fonte"], "requisito.status": "Verificado" if q["status"] == "verificado" else "Não verificado"}
            for q in g["requisitos"]]
    revisoes_arquivo = [{"arquivo.revisao": "00", "arquivo.data": "2026-10-23", "arquivo.descricao": "Emissão inicial (exemplo fictício)",
                         "arquivo.revisado_por": "Pessoa A"}]
    return {"ativos": ativos, "plano": plano, "ocorrencias": ocorrencias, "execucoes": execucoes, "falhas": falhas, "lista": lista,
            "revisoes": revisoes, "setores": setores, "pessoas": pessoas_l, "requisitos": reqs, "revisoes_arquivo": revisoes_arquivo}
